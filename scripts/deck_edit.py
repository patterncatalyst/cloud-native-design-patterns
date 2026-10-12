#!/usr/bin/env python3
"""Apply scripted edits to a .pptx in place (python-pptx + lxml; run with the S12 venv python).

    deck_edit.py <deck.pptx> <edits.json> [<edits2.json> ...] [--dry-run] [--out PATH]

Edit files are applied in order; a failing operation aborts the run and nothing is written.
Without --out the deck is rewritten in place (atomic replace) after a backup copy is saved to
$DECK_BACKUP_DIR (default ~/.cache/deck_edit/backup). --dry-run prints the text diff (using
deck_text's extractor) and writes nothing. On success a JSON summary is printed as the last line
({"ok": true, "out": ..., "created": {...}, "ops": N}).

Slides are addressed by sldId (the stable id in presentation.xml, shown as "sldId N" in the text
copy). A shape selector is a shape name (exact) or a unique substring of the shape's text; an op may
also give "shape_id": N. A sldId of the form "$name" refers to a slide created earlier by a
clone_slide op with "as": "name".

Edits JSON: a list of operations, or {"ops": [...]}.

    [
      {"op": "replace", "scope": "both", "find": "Poetry", "repl": "uv", "expect": 4,
       "sldIds": [259, 260]},                       # sldIds optional; scope slide|notes|both
      {"op": "set_notes", "sldId": 259, "paragraphs": ["First paragraph.", "Second one."]},
      {"op": "set_notes", "sldId": 259, "text": "One paragraph per line\\nof this string"},
      {"op": "set_code", "sldId": 269, "shape": "pyproject.toml",   # shape optional: sole code box
       "lines": ["# comment", "from fastapi import FastAPI", "", "app = FastAPI()"],
       "max_lines": 22, "max_cols": 95},
      {"op": "set_text", "sldId": 258, "shape": "Foundations",
       "paragraphs": ["New first paragraph", "New second paragraph"]},
      {"op": "clone_slide", "template": 269, "after": 275, "as": "r1"},   # after optional: append
      {"op": "set_text", "sldId": "$r1", "shape": "CLOUD-NATIVE", "paragraphs": ["APPENDIX R"]},
      {"op": "replace_picture", "sldId": 260, "shape": "Picture 6", "png": "/path/fig.png"},
      {"op": "set_page_number", "sldId": "$r1", "text": "273"},
      {"op": "fit_font", "sldId": 431, "shape_id": 1938, "pt": 11},
      {"op": "resize_shape", "sldId": 343, "shape_id": 979, "cy": 4572000},
      {"op": "resize_shape", "sldId": 300, "shape_id": 508, "cx": 2400000, "keep": "right"},
      {"op": "cut_code", "sldId": 474, "shape_id": 2424, "pt": 9, "cuts": [
        {"from": 13, "to": 20, "first": "# Optional defensive variant: ...", "ellipsis": "# ..."}]},
      {"op": "recolour_code", "sldId": 521, "shape_id": 8, "code_rgb": "E6E6E6", "comment_rgb": "8FB98F"}
    ]

Operations
  replace           exact text replace; "expect" (required) is the total match count over the
                    selected slides/notes and must equal the real count, otherwise the run fails.
                    Matches inside one run keep its formatting; a match spanning runs is merged
                    into the first run's formatting and logged.
  set_notes         replace the speaker notes body (creates a notes slide when missing); keeps
                    the formatting of the first existing notes paragraph.
  set_code          rewrite a code box line by line, one paragraph per line, copying the box's own
                    paragraph/run formatting; whole-line comments (# // -- /* *) take the format of
                    an existing comment line of that box. Fails over max_lines (22, or the box's current line
                    count if larger) or max_cols
                    (95) or the box capacity (height / font size / width). Optional "pt" and
                    "line_spacing" resize the box's text first, as fit_font does. A line's
                    template run is the run that put non-blank text on it. The result is then
                    coloured by line type (recolour_code): "code_rgb" (default: the box's code
                    colour, else E6E6E6) and "comment_rgb" (default: the box's comment colour,
                    else 8FB98F); "colour": false keeps the template colours.
  recolour_code     colour one code box by line type: whole-line comments "comment_rgb"
                    (8FB98F), other non-blank lines "code_rgb" (E6E6E6), a trailing-comment run
                    after code stays comment-coloured; runs of a third colour are kept unless
                    "force"; fails when nothing changes unless "allow_noop". Multi-line runs are
                    split per line first (same text).
  set_text          rewrite a text shape's paragraphs keeping per-paragraph formatting (paragraph i
                    reuses existing paragraph i, extras reuse the last one).
  clone_slide       deep copy a slide (shapes, relationships incl. images, own notes slide) and
                    insert it after another slide or at the end; the new sldId is printed and
                    returned in "created".
  replace_picture   point a picture shape at a PNG (fit inside the old frame, aspect kept).
  set_page_number   set the static page-number text box (a digits-only text at 11 pt or less).
  fit_font          set the size of every run of one shape to "pt" (a multiple of 0.5; also the
                    paragraph-end, default and bullet sizes), or multiply every size by "scale"
                    (0.5-1.0, rounded to 0.5 pt; keeps a shape's size hierarchy); optional
                    "line_spacing" (0.8-1.5, with pt); optional "space_before" (pt) resets every
                    existing paragraph space-before
                    sets every paragraph's spcPct. Floors: 9 pt for code boxes, 12 pt for other
                    text; a code box must then fit its capacity (lines and columns) or the op fails.
  cut_code          excerpt a code box in place: "cuts" is a list of {"from": a, "to": b, "first":
                    "<exact text of line a>", "last": "<line b>" (optional), "ellipsis": "# ..."
                    (optional comment line that replaces the range, coloured "comment_rgb",
                    default 8FB98F)}; 1-based line numbers of
                    the box as it is when the op runs. Every other line keeps its runs and colours.
                    Optional "pt" / "line_spacing" as in fit_font; the result must fit the box.
  resize_shape      move or resize one text shape: any of "x", "y", "cx", "cy" in EMU; "keep":
                    "right" or "bottom" (or both, as a list) holds that edge. Fails when the frame
                    would leave the slide or overlap the page-number box or the logo picture.
  set_geometry      move or resize one picture (or any shape) by "shape" name or "shape_id": any
                    of "x", "y", "w", "h" in inches; the frame must stay on the slide and clear of
                    the page-number box and the logo.

Box capacity (set_code, cut_code, fit_font): lines = 1 + (height - top inset - 1.32 x size) /
(1.20 x size x spcPct), columns = (width - insets) / (0.6 x size), the LibreOffice metrics of Red
Hat Mono measured in renders (the last line may reach into the bottom inset, not past the frame).
Run scripts/deck_fit.py on the edited deck: it checks the real render.
"""
from __future__ import annotations

import argparse
import copy
import difflib
import json
import math
import os
import re
import shutil
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402

import deck_text  # noqa: E402

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"a": A, "p": P, "r": R}
MONO_RE = deck_text.MONO_RE


def q(ns, tag):
    return f"{{{ns}}}{tag}"


class EditError(Exception):
    pass


# ----------------------------------------------------------------- helpers

def ptext(p):
    return "".join(t.text or "" for t in p.iter(q(A, "t")))


def sp_text(sp):
    return "\n".join(ptext(p) for p in sp.iter(q(A, "p")))


def all_sps(root):
    """Every p:sp in document order (groups recursed)."""
    return [e for e in root.iter(q(P, "sp"))]


def sp_name(sp):
    c = sp.find(f"{q(P, 'nvSpPr')}/{q(P, 'cNvPr')}")
    return c.get("name") if c is not None else ""


def sp_id(sp):
    c = sp.find(f"{q(P, 'nvSpPr')}/{q(P, 'cNvPr')}")
    return int(c.get("id")) if c is not None else -1


def run_font(r):
    rpr = r.find(q(A, "rPr"))
    if rpr is None:
        return None
    lat = rpr.find(q(A, "latin"))
    return lat.get("typeface") if lat is not None else None


def run_size(r):
    rpr = r.find(q(A, "rPr"))
    return int(rpr.get("sz")) if rpr is not None and rpr.get("sz") else 0


def is_code_sp(sp):
    runs = [r for r in sp.iter(q(A, "r")) if (r.findtext(q(A, "t")) or "").strip()]
    if not runs or not all(MONO_RE.search(run_font(r) or "") for r in runs):
        return False
    paras = [p for p in sp.iter(q(A, "p")) if ptext(p).strip()]
    one = len(paras) == 1 and len(ptext(paras[0])) <= 60 and max(run_size(r) for r in runs) <= 1100
    return not one  # a single short small mono line is a chip, not code


def select_sp(root, op, kind=None):
    sps = [s for s in all_sps(root) if s.find(q(P, "txBody")) is not None]
    if op.get("shape_id") is not None:
        hit = [s for s in sps if sp_id(s) == int(op["shape_id"])]
    else:
        sel = op.get("shape")
        if sel is None:
            if kind == "code":
                hit = [s for s in sps if is_code_sp(s)]
            else:
                raise EditError(f"{op['op']}: a shape selector is required")
        else:
            hit = [s for s in sps if sp_name(s) == sel]
            if not hit:
                hit = [s for s in sps if sel in sp_text(s)]
    if len(hit) != 1:
        desc = op.get("shape", op.get("shape_id", "(sole code box)"))
        names = [f"{sp_name(s)}:{sp_text(s)[:40]!r}" for s in hit]
        raise EditError(f"{op['op']}: selector {desc!r} matched {len(hit)} shapes {names}")
    return hit[0]


def set_par_text(template, text):
    """Deep-copy a paragraph, keep the first run's formatting, put `text` in it."""
    p = copy.deepcopy(template)
    runs = [c for c in p if c.tag == q(A, "r")]
    for c in list(p):
        if c.tag in (q(A, "br"), q(A, "fld")):
            p.remove(c)
    if runs:
        first = runs[0]
        for extra in runs[1:]:
            p.remove(extra)
    else:
        first = etree.SubElement(p, q(A, "r"))
        etree.SubElement(first, q(A, "t"))
        end = p.find(q(A, "endParaRPr"))
        if end is not None:
            p.remove(first)
            end.addprevious(first)
    t = first.find(q(A, "t"))
    if t is None:
        t = etree.SubElement(first, q(A, "t"))
    t.text = text
    return p


def fresh_par():
    p = etree.Element(q(A, "p"))
    r = etree.SubElement(p, q(A, "r"))
    rpr = etree.SubElement(r, q(A, "rPr"))
    rpr.set("lang", "en-US")
    etree.SubElement(r, q(A, "t"))
    return p


def rewrite_paragraphs(txbody, texts, templates=None):
    old = txbody.findall(q(A, "p"))
    base = templates or old
    if not base:
        base = [fresh_par()]
    anchor = old[0] if old else None
    new = []
    for i, text in enumerate(texts):
        new.append(set_par_text(base[min(i, len(base) - 1)], text))
    for o in old:
        txbody.remove(o)
    for p in new:
        txbody.append(p)
    return len(texts)


class Ctx:
    def __init__(self, prs):
        self.prs = prs
        self.created: dict[str, int] = {}
        self.log: list[str] = []

    def slide(self, ref):
        if isinstance(ref, str) and ref.startswith("$"):
            if ref[1:] not in self.created:
                raise EditError(f"unknown slide reference {ref}")
            ref = self.created[ref[1:]]
        for s in self.prs.slides:
            if s.slide_id == int(ref):
                return s
        raise EditError(f"no slide with sldId {ref}")

    def sld_id_list(self):
        return self.prs.slides._sldIdLst


# ----------------------------------------------------------------- replace

def replace_in_paragraph(p, find, repl, label, log):
    """Replace in one paragraph; returns match count."""
    count = 0
    # segments of consecutive a:r separated by a:br
    segs, cur = [], []
    for c in p:
        if c.tag == q(A, "r"):
            cur.append(c)
        elif c.tag == q(A, "br"):
            segs.append(cur)
            cur = []
    segs.append(cur)
    for runs in segs:
        if not runs:
            continue
        texts = [(r.findtext(q(A, "t")) or "") for r in runs]
        full = "".join(texts)
        spans, pos = [], 0
        while True:
            i = full.find(find, pos)
            if i < 0:
                break
            spans.append((i, i + len(find)))
            pos = i + len(find)
        if not spans:
            continue
        count += len(spans)
        bounds, acc = [], 0
        for t in texts:
            bounds.append((acc, acc + len(t)))
            acc += len(t)

        def run_at(off, end=False):
            for k, (s, e) in enumerate(bounds):
                if (s < off <= e) if end else (s <= off < e):
                    return k
            return len(bounds) - 1

        for s, e in reversed(spans):
            i, j = run_at(s), run_at(e, end=True)
            if i == j:
                bs = bounds[i][0]
                texts[i] = texts[i][: s - bs] + repl + texts[i][e - bs:]
            else:
                log.append(f"cross-run match in {label}: {find!r} spans runs {i}..{j}; merged into run {i} formatting")
                bs_i, bs_j = bounds[i][0], bounds[j][0]
                texts[i] = texts[i][: s - bs_i] + repl
                for k in range(i + 1, j):
                    texts[k] = ""
                texts[j] = texts[j][e - bs_j:]
        for r, t in zip(runs, texts):
            te = r.find(q(A, "t"))
            if te is None:
                te = etree.SubElement(r, q(A, "t"))
            te.text = t
        # drop runs emptied by a cross-run merge (keep the first run of the segment)
        for k, r in enumerate(runs):
            if k > 0 and (r.findtext(q(A, "t")) or "") == "" and bounds[k][1] > bounds[k][0]:
                p.remove(r)
    return count


def notes_roots(slide):
    if not slide.has_notes_slide:
        return []
    return [slide.notes_slide._element]


def op_replace(ctx, op):
    for k in ("find", "repl", "expect"):
        if k not in op:
            raise EditError(f"replace: missing {k!r}")
    find, repl, expect = op["find"], op["repl"], int(op["expect"])
    if not find:
        raise EditError("replace: empty find")
    scope = op.get("scope", "both")
    if scope not in ("slide", "notes", "both"):
        raise EditError(f"replace: bad scope {scope!r}")
    ids = op.get("sldIds")
    slides = [ctx.slide(i) for i in ids] if ids is not None else list(ctx.prs.slides)
    total = 0
    for s in slides:
        roots = []
        if scope in ("slide", "both"):
            roots.append(("slide", s._element))
        if scope in ("notes", "both"):
            for r in notes_roots(s):
                # only the notes body placeholder text, never the slide-number or image placeholders
                for sp in all_sps(r):
                    ph = sp.find(f"{q(P, 'nvSpPr')}/{q(P, 'nvPr')}/{q(P, 'ph')}")
                    if ph is not None and ph.get("type") == "body":
                        roots.append(("notes", sp))
        for kind, root in roots:
            for p in root.iter(q(A, "p")):
                total += replace_in_paragraph(p, find, repl, f"sldId {s.slide_id} {kind}", ctx.log)
    if total != expect:
        raise EditError(f"replace {find!r}: expected {expect} match(es), found {total}")
    ctx.log.append(f"replace {find!r} -> {repl!r}: {total} match(es)")


# ----------------------------------------------------------------- notes

def notes_body(slide, create=True):
    if not slide.has_notes_slide and not create:
        return None
    ns = slide.notes_slide
    for sp in all_sps(ns._element):
        ph = sp.find(f"{q(P, 'nvSpPr')}/{q(P, 'nvPr')}/{q(P, 'ph')}")
        if ph is not None and ph.get("type") == "body":
            return sp
    raise EditError(f"sldId {slide.slide_id}: notes slide has no body placeholder")


def op_set_notes(ctx, op):
    s = ctx.slide(op["sldId"])
    if "paragraphs" in op:
        texts = list(op["paragraphs"])
    elif "text" in op:
        texts = op["text"].split("\n")
    else:
        raise EditError("set_notes: needs text or paragraphs")
    texts = [t for t in texts if t.strip() != ""] or [""]
    existed = s.has_notes_slide
    sp = notes_body(s)
    tx = sp.find(q(P, "txBody"))
    if tx is None:
        raise EditError(f"sldId {s.slide_id}: notes body has no txBody")
    old = [p for p in tx.findall(q(A, "p")) if ptext(p).strip()] or tx.findall(q(A, "p"))
    rewrite_paragraphs(tx, texts, templates=old[:1] or None)
    ctx.log.append(f"set_notes sldId {s.slide_id}: {len(texts)} paragraph(s)" + ("" if existed else " (created notes slide)"))


# ----------------------------------------------------------------- code

COMMENT_RE = re.compile(r"^\s*(#(?!\s*(include|define|pragma|if|ifdef|ifndef|else|elif|endif|undef|error)\b)|#!|//|--\s|/\*|\*\s|\*/)")


def is_comment(text):
    return bool(COMMENT_RE.match(text)) and not text.lstrip().startswith("#!")


# LibreOffice line pitch for Red Hat Mono and Red Hat Text: 1.20 x the font size x the
# paragraph's spcPct (median of 248 multi-line text boxes measured in the soffice PDFs of the
# Python and .NET decks with pdftotext -bbox-layout, 2026-10-12; range 1.1985-1.2020).
# Red Hat Mono advances 0.5985 x the size per character; a rendered line box is 1.32 x the size.
LINE_HEIGHT = 1.20
MONO_ADVANCE = 0.6
GLYPH_BOX = 1.32   # ascent + descent of a rendered line box (pdftotext -bbox-layout), x the size


def box_capacity(sp, tmpl, run=None):
    xfrm = sp.find(f"{q(P, 'spPr')}/{q(A, 'xfrm')}")
    ext = xfrm.find(q(A, "ext")) if xfrm is not None else None
    if ext is None:
        return None, None
    body = sp.find(f"{q(P, 'txBody')}/{q(A, 'bodyPr')}")
    ti = int(body.get("tIns", 45720)) if body is not None else 45720
    bi = int(body.get("bIns", 45720)) if body is not None else 45720
    li = int(body.get("lIns", 91440)) if body is not None else 91440
    ri = int(body.get("rIns", 91440)) if body is not None else 91440
    run = run if run is not None else tmpl.find(q(A, "r"))
    sz = (run_size(run) if run is not None else 0) or 1300
    sp_pct = tmpl.find(f"{q(A, 'pPr')}/{q(A, 'lnSpc')}/{q(A, 'spcPct')}")
    spc = int(sp_pct.get("val")) / 100000 if sp_pct is not None else 1.0
    line_pt = sz / 100 * LINE_HEIGHT * spc
    # n lines end at top inset + (n - 1) pitches + the last line's glyph box (GLYPH_BOX x size);
    # that box may reach into the bottom inset, as long as it stays inside the frame
    lines = math.floor(((int(ext.get("cy")) - ti) / 12700 - sz / 100 * GLYPH_BOX) / line_pt + 1 + 0.02)
    cols = math.floor((int(ext.get("cx")) - li - ri) / 12700 / (sz / 100 * MONO_ADVANCE))
    return lines, cols


def code_lines(tx):
    """Existing lines of a code box as dicts {p, run, br, text}; handles one paragraph per line,
    runs separated by a:br inside paragraphs, and newlines embedded in a:t."""
    out = []
    for p in tx.findall(q(A, "p")):
        cur = {"p": p, "run": None, "br": None, "text": "", "nonblank": False}
        for c in list(p):
            if c.tag == q(A, "r"):
                t = c.findtext(q(A, "t")) or ""
                parts = t.split("\n")
                for k, part in enumerate(parts):
                    if k > 0:
                        out.append(cur)
                        cur = {"p": p, "run": None, "br": None, "text": "", "nonblank": False}
                    cur["text"] += part
                    # the line's template run is the first run that put non-blank text on this
                    # line; a run that only ends here (e.g. "# comment\n" before a code run)
                    # is a fallback for a blank line, never the template of a code line
                    if part.strip():
                        if not cur["nonblank"]:
                            cur["run"] = c
                            cur["nonblank"] = True
                    elif cur["run"] is None:
                        cur["run"] = c
            elif c.tag == q(A, "br"):
                cur["br"] = c
                out.append(cur)
                cur = {"p": p, "run": None, "br": None, "text": "", "nonblank": False}
        out.append(cur)
    return out


# House code palette: code lines E6E6E6, whole-line comments 8FB98F (the 8A8A8A grey is the chip
# label colour and never appears in a code box).
CODE_RGB = "E6E6E6"
COMMENT_RGB = "8FB98F"
RGB_RE = re.compile(r"^[0-9A-Fa-f]{6}$")


def run_rgb(r):
    rpr = r.find(q(A, "rPr"))
    sf = rpr.find(q(A, "solidFill")) if rpr is not None else None
    clr = sf.find(q(A, "srgbClr")) if sf is not None else None
    return clr.get("val").upper() if clr is not None and clr.get("val") else None


def set_run_rgb(r, rgb):
    rpr = r.find(q(A, "rPr"))
    if rpr is None:
        rpr = etree.Element(q(A, "rPr"))
        r.insert(0, rpr)
    for tag in ("noFill", "gradFill", "blipFill", "pattFill", "grpFill"):
        for e in rpr.findall(q(A, tag)):
            rpr.remove(e)
    sf = rpr.find(q(A, "solidFill"))
    if sf is None:
        sf = etree.Element(q(A, "solidFill"))
        ln = rpr.find(q(A, "ln"))
        if ln is not None:
            ln.addnext(sf)
        else:
            rpr.insert(0, sf)
    for c in list(sf):
        sf.remove(c)
    etree.SubElement(sf, q(A, "srgbClr")).set("val", rgb)


def check_rgb(v, label):
    if v is not None and not RGB_RE.match(str(v)):
        raise EditError(f"{label}: {v!r} is not an RRGGBB colour")
    return v.upper() if v is not None else None


def _split_multiline_runs(tx):
    """Split every run whose text holds newlines into one run per line (each piece but the last
    keeps its trailing newline), so a line's colour can be set on its own; the text is unchanged."""
    n = 0
    for r in list(tx.iter(q(A, "r"))):
        t = r.find(q(A, "t"))
        text = t.text if t is not None and t.text else ""
        if "\n" not in text.rstrip("\n") and text.count("\n") <= 1:
            continue
        pieces = [x + "\n" for x in text.split("\n")[:-1]] + [text.split("\n")[-1]]
        pieces = [x for x in pieces if x]
        if len(pieces) <= 1:
            continue
        t.text = pieces[0]
        prev = r
        for piece in pieces[1:]:
            nr = copy.deepcopy(r)
            nr.find(q(A, "t")).text = piece
            prev.addnext(nr)
            prev = nr
        n += 1
    return n


def recolour_box(tx, code_rgb=CODE_RGB, comment_rgb=COMMENT_RGB, force=False):
    """Colour a code box by line type: whole-line comments comment_rgb, every other non-blank
    line code_rgb; a run that starts a trailing comment after code on the same line keeps
    comment_rgb. Runs of any third colour are left alone unless force. Returns the number of
    runs whose colour changed."""
    _split_multiline_runs(tx)
    lines, cur = [], []
    for p in tx.findall(q(A, "p")):
        cur = []
        for c in p:
            if c.tag == q(A, "r"):
                text = c.findtext(q(A, "t")) or ""
                cur.append((c, text.rstrip("\n")))
                if text.endswith("\n"):
                    lines.append(cur)
                    cur = []
            elif c.tag == q(A, "br"):
                lines.append(cur)
                cur = []
        lines.append(cur)
    ours = {code_rgb, comment_rgb, CODE_RGB, COMMENT_RGB, None}
    changed = 0
    for segs in lines:
        text = "".join(t for _, t in segs)
        if not text.strip():
            continue
        whole = is_comment(text)
        seen_code, in_comment = False, False
        for r, t in segs:
            if not t.strip():
                continue
            if whole:
                want = comment_rgb
            else:
                if seen_code and not in_comment and is_comment(t):
                    in_comment = True
                want = comment_rgb if in_comment else code_rgb
                if not is_comment(t):
                    seen_code = True
            have = run_rgb(r)
            if have == want or (have not in ours and not force):
                continue
            set_run_rgb(r, want)
            changed += 1
    return changed


def op_recolour_code(ctx, op):
    """Recolour one code box by line type (see recolour_box)."""
    s = ctx.slide(op["sldId"])
    sp = select_sp(s._element, op, kind="code")
    label = f"recolour_code sldId {s.slide_id}"
    code_rgb = check_rgb(op.get("code_rgb", CODE_RGB), label)
    comment_rgb = check_rgb(op.get("comment_rgb", COMMENT_RGB), label)
    n = recolour_box(sp.find(q(P, "txBody")), code_rgb, comment_rgb, bool(op.get("force")))
    if n == 0 and not op.get("allow_noop"):
        raise EditError(f"{label}: no run changed colour (wrong shape, or already coloured)")
    ctx.log.append(f"{label} {sp_name(sp)!r}: {n} run(s) recoloured (code {code_rgb}, comment {comment_rgb})")


def op_set_code(ctx, op):
    s = ctx.slide(op["sldId"])
    sp = select_sp(s._element, op, kind="code")
    tx = sp.find(q(P, "txBody"))
    if "lines" in op:
        lines = list(op["lines"])
    elif "code" in op:
        lines = op["code"].rstrip("\n").split("\n")
    else:
        raise EditError("set_code: needs lines or code")
    lines = [ln.replace("\t", "    ").rstrip() for ln in lines]
    if op.get("pt") is not None or op.get("line_spacing") is not None:
        if op.get("pt") is None:
            raise EditError(f"set_code sldId {s.slide_id}: line_spacing needs pt")
        size_shape(sp, op["pt"], op.get("line_spacing"), f"set_code sldId {s.slide_id}", code=True)
    old_lines = code_lines(tx)
    has_br = any(l["br"] is not None for l in old_lines)
    embedded = any("\n" in (t.text or "") for t in tx.iter(q(A, "t")))
    br_model = has_br or (len(tx.findall(q(A, "p"))) == 1 and embedded)
    code_l = comment_l = blank_l = None
    for l in old_lines:
        if l["run"] is None and not l["text"].strip():
            blank_l = blank_l or l
        elif not l["text"].strip():
            blank_l = blank_l or l
        elif is_comment(l["text"]):
            comment_l = comment_l or l
        else:
            code_l = code_l or l
    if code_l is None:
        code_l = comment_l
    if code_l is None or code_l["run"] is None:
        raise EditError(f"set_code sldId {s.slide_id}: the box has no text line to copy formatting from")
    if comment_l is None:
        comment_l = code_l
        ctx.log.append(f"set_code sldId {s.slide_id}: no comment line in the box; comments use the code format")
    # Default limit: 22 lines, or the box's current line count when it already holds more
    # (the existing font size was chosen for that many lines).
    max_lines = int(op.get("max_lines", max(22, len(old_lines))))
    max_cols = int(op.get("max_cols", 95))
    cap_l, cap_c = box_capacity(sp, code_l["p"], code_l["run"])
    if len(lines) > max_lines:
        raise EditError(f"set_code sldId {s.slide_id}: {len(lines)} lines exceeds max_lines {max_lines}")
    if cap_l is not None and len(lines) > cap_l:
        raise EditError(f"set_code sldId {s.slide_id}: {len(lines)} lines exceeds the box capacity {cap_l}")
    wide = max((len(ln) for ln in lines), default=0)
    if wide > max_cols:
        raise EditError(f"set_code sldId {s.slide_id}: a line is {wide} columns, max_cols {max_cols}")
    if cap_c is not None and wide > cap_c:
        raise EditError(f"set_code sldId {s.slide_id}: a line is {wide} columns, box capacity {cap_c}")
    if not br_model:
        old = tx.findall(q(A, "p"))
        new = []
        for ln in lines:
            if not ln.strip():
                base = blank_l["p"] if blank_l is not None else code_l["p"]
                new.append(set_par_text(base, ""))
            else:
                base = (comment_l if is_comment(ln) else code_l)["p"]
                new.append(set_par_text(base, ln))
        for o in old:
            tx.remove(o)
        for p in new:
            tx.append(p)
    else:
        p0 = copy.deepcopy(old_lines[0]["p"])
        end = p0.find(q(A, "endParaRPr"))
        for c in list(p0):
            if c.tag in (q(A, "r"), q(A, "br"), q(A, "fld")):
                p0.remove(c)

        def add(el):
            if end is not None:
                end.addprevious(el)
            else:
                p0.append(el)

        def br_for(tmpl_line):
            if tmpl_line["br"] is not None:
                return copy.deepcopy(tmpl_line["br"])
            b = etree.Element(q(A, "br"))
            rpr = tmpl_line["run"].find(q(A, "rPr"))
            if rpr is not None:
                b.append(copy.deepcopy(rpr))
            return b

        for i, ln in enumerate(lines):
            tmpl = comment_l if is_comment(ln) else code_l
            if ln.strip():
                r = copy.deepcopy(tmpl["run"])
                for extra in r.findall(q(A, "t"))[1:]:
                    r.remove(extra)
                te = r.find(q(A, "t"))
                if te is None:
                    te = etree.SubElement(r, q(A, "t"))
                te.text = ln
                add(r)
            if i < len(lines) - 1:
                add(br_for(tmpl if ln.strip() else code_l))
        for o in tx.findall(q(A, "p")):
            tx.remove(o)
        tx.append(p0)
    recoloured = 0
    if op.get("colour", True):
        label = f"set_code sldId {s.slide_id}"
        have_code = run_rgb(code_l["run"])
        have_comment = run_rgb(comment_l["run"]) if comment_l is not code_l else None
        code_rgb = check_rgb(op.get("code_rgb"), label) or (
            have_code if have_code and have_code != COMMENT_RGB else CODE_RGB)
        comment_rgb = check_rgb(op.get("comment_rgb"), label) or (
            have_comment if have_comment and have_comment != code_rgb else COMMENT_RGB)
        recoloured = recolour_box(tx, code_rgb, comment_rgb)
    ctx.log.append(f"set_code sldId {s.slide_id}: {len(lines)} line(s), widest {wide} col(s) "
                   f"(box {cap_l} lines x {cap_c} cols, {'br' if br_model else 'paragraph'} model)"
                   + (f", {recoloured} run(s) recoloured" if recoloured else ""))


def op_set_text(ctx, op):
    s = ctx.slide(op["sldId"])
    sp = select_sp(s._element, op)
    tx = sp.find(q(P, "txBody"))
    texts = list(op["paragraphs"])
    templates = [p for p in tx.findall(q(A, "p")) if ptext(p).strip()] or tx.findall(q(A, "p"))
    rewrite_paragraphs(tx, texts, templates=templates)
    ctx.log.append(f"set_text sldId {s.slide_id} {sp_name(sp)!r}: {len(texts)} paragraph(s)")


def op_set_page_number(ctx, op):
    s = ctx.slide(op["sldId"])
    hits = []
    for sp in all_sps(s._element):
        if sp.find(q(P, "txBody")) is None:
            continue
        t = sp_text(sp).strip()
        runs = [r for r in sp.iter(q(A, "r"))]
        size = max([run_size(r) for r in runs] or [0])
        if re.fullmatch(r"\d{1,3}", t) and 0 < size <= 1100:
            hits.append(sp)
    if len(hits) != 1:
        raise EditError(f"set_page_number sldId {s.slide_id}: found {len(hits)} page-number boxes")
    r = next(hits[0].iter(q(A, "r")))
    r.find(q(A, "t")).text = str(op["text"])
    ctx.log.append(f"set_page_number sldId {s.slide_id}: {op['text']}")


# ----------------------------------------------------------------- picture

def op_replace_picture(ctx, op):
    s = ctx.slide(op["sldId"])
    pics = list(s._element.iter(q(P, "pic")))
    sel = op.get("shape")
    if op.get("shape_id") is not None:
        hit = [p for p in pics if p.find(f"{q(P, 'nvPicPr')}/{q(P, 'cNvPr')}").get("id") == str(op["shape_id"])]
    elif sel is None and len(pics) == 1:
        hit = pics
    else:
        hit = [p for p in pics if p.find(f"{q(P, 'nvPicPr')}/{q(P, 'cNvPr')}").get("name") == sel]
        if not hit:
            hit = [p for p in pics if sel in (p.find(f"{q(P, 'nvPicPr')}/{q(P, 'cNvPr')}").get("descr") or "")]
    if len(hit) != 1:
        raise EditError(f"replace_picture sldId {s.slide_id}: selector matched {len(hit)} pictures of {len(pics)}")
    pic = hit[0]
    png = Path(op["png"])
    if not png.is_file():
        raise EditError(f"replace_picture: {png} not found")
    blip = pic.find(f".//{q(A, 'blip')}")
    old_rid = blip.get(q(R, "embed"))
    img_part, new_rid = s.part.get_or_add_image_part(str(png))
    blip.set(q(R, "embed"), new_rid)
    if old_rid != new_rid:
        # drop the old relationship when nothing else on the slide uses it
        still = any(v == old_rid for e in s._element.iter() for k, v in e.attrib.items() if k.startswith("{" + R))
        if not still:
            s.part.drop_rel(old_rid)
    if op.get("fit", True):
        iw, ih = img_part.image.size
        xfrm = pic.find(f"{q(P, 'spPr')}/{q(A, 'xfrm')}")
        off, ext = xfrm.find(q(A, "off")), xfrm.find(q(A, "ext"))
        bx, by, bw, bh = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))
        scale = min(bw / iw, bh / ih)
        nw, nh = int(iw * scale), int(ih * scale)
        off.set("x", str(bx + (bw - nw) // 2))
        off.set("y", str(by + (bh - nh) // 2))
        ext.set("cx", str(nw))
        ext.set("cy", str(nh))
        sr = pic.find(f"{q(P, 'blipFill')}/{q(A, 'srcRect')}")
        if sr is not None:
            sr.getparent().remove(sr)
    ctx.log.append(f"replace_picture sldId {s.slide_id}: {png.name}")


# ----------------------------------------------------------------- clone

def op_clone_slide(ctx, op):
    prs = ctx.prs
    tpl = ctx.slide(op["template"])
    after = ctx.slide(op["after"]) if op.get("after") is not None else None
    new = prs.slides.add_slide(tpl.slide_layout)
    # a free partname (python-pptx numbers by slide count, which collides after deletions)
    pkg_names = {str(p.partname) for p in prs.part.package.iter_parts()}
    if sum(1 for n in pkg_names if n == str(new.part.partname)) > 1:
        n = 1
        while f"/ppt/slides/slide{n}.xml" in pkg_names:
            n += 1
        new.part.partname = type(new.part.partname)(f"/ppt/slides/slide{n}.xml")
    # relationships: same targets, mapped rIds
    rid_map = {}
    for rid, rel in tpl.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    # replace the XML children wholesale with copies of the template's
    for c in list(new._element):
        new._element.remove(c)
    for c in tpl._element:
        new._element.append(copy.deepcopy(c))
    for e in new._element.iter():
        for k, v in list(e.attrib.items()):
            if k.startswith("{" + R) and v in rid_map:
                e.set(k, rid_map[v])
            elif k.startswith("{" + R):
                raise EditError(f"clone_slide: unmapped relationship {v} on {e.tag}")
    # notes: own notes slide, content copied from the template's
    if tpl.has_notes_slide:
        nn = new.notes_slide
        src = tpl.notes_slide._element
        dst = nn._element
        dst_csld, src_csld = dst.find(q(P, "cSld")), src.find(q(P, "cSld"))
        dst.replace(dst_csld, copy.deepcopy(src_csld))
        for e in dst.iter():
            for k in e.attrib:
                if k.startswith("{" + R):
                    raise EditError("clone_slide: notes slide holds a relationship reference; not supported")
    # position
    lst = ctx.sld_id_list()
    new_el = next(e for e in lst if int(e.get("id")) == new.slide_id)
    lst.remove(new_el)
    if after is None:
        lst.append(new_el)
    else:
        tgt = next(e for e in lst if int(e.get("id")) == after.slide_id)
        tgt.addnext(new_el)
    name = op.get("as")
    if name:
        ctx.created[name] = new.slide_id
    ctx.log.append(f"clone_slide {tpl.slide_id} -> new sldId {new.slide_id}"
                   + (f" after {after.slide_id}" if after else " at end") + (f" as ${name}" if name else ""))


# ----------------------------------------------------------------- fit

MIN_CODE_PT = 9
MIN_TEXT_PT = 12


def size_shape(sp, pt, line_spacing, label, code=None):
    """Set every run, paragraph-end, default and bullet size of a shape to pt (and optionally
    every paragraph's line spacing); enforces the floors (code: set_code and cut_code boxes, or a
    shape whose runs are all monospace)."""
    code = is_code_sp(sp) if code is None else code
    pt = float(pt)
    floor = MIN_CODE_PT if code else MIN_TEXT_PT
    if pt < floor:
        raise EditError(f"{label}: {pt} pt is under the {floor} pt floor for {'code' if code else 'text'}")
    if pt * 2 != int(pt * 2):
        raise EditError(f"{label}: {pt} pt is not a multiple of 0.5")
    tx = sp.find(q(P, "txBody"))
    val = str(int(round(pt * 100)))
    n = 0
    # a paragraph without runs takes its height from endParaRPr (or the master default when
    # that has no size), so every paragraph gets one with the new size
    for p in tx.findall(q(A, "p")):
        if p.find(q(A, "endParaRPr")) is None:
            etree.SubElement(p, q(A, "endParaRPr")).set("lang", "en-US")
    for tag in ("rPr", "endParaRPr", "defRPr"):
        for e in tx.iter(q(A, tag)):
            e.set("sz", val)
            n += 1
    for e in tx.iter(q(A, "buSzPts")):
        e.set("val", val)
    if n == 0:
        raise EditError(f"{label}: the shape has no run properties to size")
    if line_spacing is not None:
        pct = int(round(float(line_spacing) * 100000))
        if not 80000 <= pct <= 150000:
            raise EditError(f"{label}: line_spacing {line_spacing} outside 0.8-1.5")
        for p in tx.findall(q(A, "p")):
            ppr = p.find(q(A, "pPr"))
            if ppr is None:
                ppr = etree.Element(q(A, "pPr"))
                p.insert(0, ppr)
            ln = ppr.find(q(A, "lnSpc"))
            if ln is not None:
                ppr.remove(ln)
            ln = etree.Element(q(A, "lnSpc"))
            etree.SubElement(ln, q(A, "spcPct")).set("val", str(pct))
            ppr.insert(0, ln)
    return code


def scale_shape(sp, factor, label):
    """Multiply every size in a shape by factor (rounded to 0.5 pt); keeps the ratio between a
    shape's sizes. The smallest resulting text size must respect the floor."""
    code = is_code_sp(sp)
    floor = MIN_CODE_PT if code else MIN_TEXT_PT
    if not 0.5 <= factor < 1.0:
        raise EditError(f"{label}: scale {factor} outside 0.5-1.0")
    tx = sp.find(q(P, "txBody"))
    run_sizes = [run_size(r) for r in tx.iter(q(A, "r")) if run_size(r)]
    if not run_sizes:
        raise EditError(f"{label}: the shape has no sized runs")
    new = lambda v: int(round(int(v) * factor / 50.0)) * 50
    if new(min(run_sizes)) < floor * 100:
        raise EditError(f"{label}: scale {factor} takes {min(run_sizes) / 100} pt under the {floor} pt floor")
    for tag in ("rPr", "endParaRPr", "defRPr"):
        for e in tx.iter(q(A, tag)):
            if e.get("sz"):
                e.set("sz", str(new(e.get("sz"))))
    for e in tx.iter(q(A, "buSzPts")):
        e.set("val", str(new(e.get("val"))))
    return code


def op_fit_font(ctx, op):
    """Set the font size of every run (and the paragraph-end and bullet sizes) of one shape, or
    scale all of them by a factor."""
    s = ctx.slide(op["sldId"])
    sp = select_sp(s._element, op)
    if ("pt" in op) == ("scale" in op):
        raise EditError("fit_font: give exactly one of pt or scale")
    if "scale" in op:
        code = scale_shape(sp, float(op["scale"]), f"fit_font sldId {s.slide_id}")
        pt = f"x{op['scale']}"
    else:
        pt = float(op["pt"])
        code = size_shape(sp, pt, op.get("line_spacing"), f"fit_font sldId {s.slide_id}")
    tx = sp.find(q(P, "txBody"))
    if op.get("space_before") is not None:
        sb = float(op["space_before"])
        if not 0 <= sb <= 24:
            raise EditError(f"fit_font sldId {s.slide_id}: space_before {sb} outside 0-24 pt")
        hits = 0
        for e in tx.iter(q(A, "spcBef")):
            for c in list(e):
                e.remove(c)
            etree.SubElement(e, q(A, "spcPts")).set("val", str(int(round(sb * 100))))
            hits += 1
        if not hits:
            raise EditError(f"fit_font sldId {s.slide_id}: no paragraph spacing to change")
    note = ""
    if code:
        lines = code_lines(tx)
        tmpl = next((l for l in lines if l["run"] is not None and l["text"].strip()), None)
        if tmpl is not None:
            cap_l, cap_c = box_capacity(sp, tmpl["p"], tmpl["run"])
            wide = max(len(l["text"].replace("\t", "    ").rstrip()) for l in lines)
            if cap_l is not None and len(lines) > cap_l:
                raise EditError(f"fit_font sldId {s.slide_id}: {len(lines)} lines still exceed the box "
                                f"capacity {cap_l} at {pt} pt")
            if cap_c is not None and wide > cap_c:
                raise EditError(f"fit_font sldId {s.slide_id}: a {wide}-column line still exceeds the "
                                f"box capacity {cap_c} at {pt} pt")
            note = f" (box {cap_l} lines x {cap_c} cols, {len(lines)} lines, widest {wide})"
    ctx.log.append(f"fit_font sldId {s.slide_id} {sp_name(sp)!r}: {pt} pt{note}")


def _line_records(tx):
    """Lines of a code box in order: {"p", "elems", "br", "text"}; a line ends at an a:br or at
    the end of its paragraph. Newlines embedded in a:t are not supported here."""
    recs = []
    for p in tx.findall(q(A, "p")):
        cur = {"p": p, "elems": [], "br": None, "text": ""}
        for c in p:
            if c.tag in (q(A, "r"), q(A, "fld")):
                t = "".join(x.text or "" for x in c.iter(q(A, "t")))
                if "\n" in t:
                    raise EditError("cut_code: newlines inside a text run are not supported")
                cur["elems"].append(c)
                cur["text"] += t
            elif c.tag == q(A, "br"):
                cur["br"] = c
                recs.append(cur)
                cur = {"p": p, "elems": [], "br": None, "text": ""}
        recs.append(cur)
    return recs


def op_cut_code(ctx, op):
    """Excerpt a code box in place: remove line ranges, optionally replacing each with one comment
    line (the ellipsis), keeping the formatting of every remaining line."""
    s = ctx.slide(op["sldId"])
    sp = select_sp(s._element, op, kind="code")
    tx = sp.find(q(P, "txBody"))
    label = f"cut_code sldId {s.slide_id}"
    if op.get("pt") is not None:
        size_shape(sp, op["pt"], op.get("line_spacing"), label, code=True)
    recs = _line_records(tx)
    cuts = sorted(op.get("cuts") or [], key=lambda c: int(c["from"]))
    if not cuts:
        raise EditError(f"{label}: no cuts")
    removed, ellipsis_at = set(), {}
    last_end = 0
    for c in cuts:
        a, b = int(c["from"]), int(c.get("to", c["from"]))
        if a <= last_end or b < a or b > len(recs):
            raise EditError(f"{label}: bad or overlapping range {a}-{b} (box has {len(recs)} lines)")
        last_end = b
        if recs[a - 1]["text"].rstrip() != c["first"].rstrip():
            raise EditError(f"{label}: line {a} is {recs[a - 1]['text']!r}, expected {c['first']!r}")
        if "last" in c and recs[b - 1]["text"].rstrip() != c["last"].rstrip():
            raise EditError(f"{label}: line {b} is {recs[b - 1]['text']!r}, expected {c['last']!r}")
        removed.update(range(a - 1, b))
        if c.get("ellipsis"):
            if not is_comment(c["ellipsis"]):
                raise EditError(f"{label}: ellipsis {c['ellipsis']!r} is not a comment line")
            ellipsis_at[a - 1] = c["ellipsis"]
    comment = next((r for r in recs if r["text"].strip() and is_comment(r["text"]) and r["elems"]), None)
    code = next((r for r in recs if r["text"].strip() and not is_comment(r["text"]) and r["elems"]), None)
    tmpl = comment or code
    if ellipsis_at and tmpl is None:
        raise EditError(f"{label}: no line to copy the ellipsis format from")
    if ellipsis_at and comment is None:
        ctx.log.append(f"{label}: no comment line in the box; the ellipsis uses the code format")
    br_tmpl = next((r["br"] for r in recs if r["br"] is not None), None)

    def ellipsis_run(text):
        r = copy.deepcopy(next(e for e in tmpl["elems"] if e.tag == q(A, "r") and (e.findtext(q(A, "t")) or "").strip()))
        for extra in r.findall(q(A, "t"))[1:]:
            r.remove(extra)
        r.find(q(A, "t")).text = text
        # an ellipsis is a comment line: comment colour even when the template is a code line
        set_run_rgb(r, ell_rgb)
        return r

    ell_rgb = check_rgb(op.get("comment_rgb"), label) or COMMENT_RGB
    paras = tx.findall(q(A, "p"))
    by_p = {id(p): [i for i, r in enumerate(recs) if r["p"] is p] for p in paras}
    kept_total = 0
    for p in paras:
        idx = by_p[id(p)]
        new = []
        for i in idx:
            if i in ellipsis_at:
                new.append(("ell", ellipsis_at[i]))
            if i not in removed:
                new.append(("rec", recs[i]))
        kept_total += len(new)
        if new == [("rec", recs[idx[0]])] and len(idx) == 1:
            continue  # untouched single-line paragraph
        if len(idx) == 1 and not new:
            tx.remove(p)
            continue
        if len(idx) == 1 and len(new) == 1 and new[0][0] == "ell":
            np_ = set_par_text(tmpl["p"] if len(by_p[id(tmpl["p"])]) == 1 else p, new[0][1])
            if len(by_p[id(tmpl["p"])]) != 1:
                # template line sits in a multi-line paragraph: rebuild this one from the run
                for c in list(np_):
                    if c.tag in (q(A, "r"), q(A, "fld"), q(A, "br")):
                        np_.remove(c)
                end = np_.find(q(A, "endParaRPr"))
                (end.addprevious if end is not None else np_.append)(ellipsis_run(new[0][1]))
            for r in np_.findall(q(A, "r")):
                set_run_rgb(r, ell_rgb)
            p.addprevious(np_)
            tx.remove(p)
            continue
        # multi-line paragraph (a:br model): rebuild its runs and breaks
        if not new:
            tx.remove(p)
            continue
        end = p.find(q(A, "endParaRPr"))
        for c in list(p):
            if c.tag in (q(A, "r"), q(A, "fld"), q(A, "br")):
                p.remove(c)
        add = end.addprevious if end is not None else p.append
        for k, (kind, v) in enumerate(new):
            if kind == "ell":
                add(ellipsis_run(v))
            else:
                for e in v["elems"]:
                    add(e)
            if k < len(new) - 1:
                b = v["br"] if kind == "rec" and v["br"] is not None else None
                if b is None:
                    if br_tmpl is None:
                        b = etree.Element(q(A, "br"))
                    else:
                        b = copy.deepcopy(br_tmpl)
                add(b)
    if not tx.findall(q(A, "p")):
        raise EditError(f"{label}: every line was cut")
    lines = code_lines(tx)
    first = next((l for l in lines if l["run"] is not None and l["text"].strip()), None)
    cap_l, cap_c = box_capacity(sp, first["p"], first["run"]) if first is not None else (None, None)
    wide = max(len(l["text"].replace("\t", "    ").rstrip()) for l in lines)
    if cap_l is not None and len(lines) > cap_l:
        raise EditError(f"{label}: {len(lines)} lines after the cuts exceed the box capacity {cap_l}")
    if cap_c is not None and wide > cap_c:
        raise EditError(f"{label}: a {wide}-column line exceeds the box capacity {cap_c}")
    ctx.log.append(f"{label}: {len(recs)} -> {len(lines)} line(s), {len(ellipsis_at)} ellipsis line(s)"
                   f" (box {cap_l} lines x {cap_c} cols)")


def _furniture(slide_el, sw, sh):
    """Frames (x0, y0, x1, y1, what, element) of the page-number box and the logo picture."""
    out = []
    for e in slide_el.iter(q(P, "sp"), q(P, "pic")):
        xf = e.find(f"{q(P, 'spPr')}/{q(A, 'xfrm')}")
        if xf is None or xf.find(q(A, "off")) is None or xf.find(q(A, "ext")) is None:
            continue
        x, y = int(xf.find(q(A, "off")).get("x")), int(xf.find(q(A, "off")).get("y"))
        cx, cy = int(xf.find(q(A, "ext")).get("cx")), int(xf.find(q(A, "ext")).get("cy"))
        if e.tag == q(P, "pic"):
            if y >= sh * 0.85 and cx < sw * 0.25:
                out.append((x, y, x + cx, y + cy, "logo", e))
        elif e.find(q(P, "txBody")) is not None:
            t = sp_text(e).strip()
            size = max([run_size(r) for r in e.iter(q(A, "r"))] or [0])
            if re.fullmatch(r"\d{1,3}", t) and 0 < size <= 1100:
                out.append((x, y, x + cx, y + cy, "page number", e))
    return out


def op_resize_shape(ctx, op):
    """Move or resize one text shape: any of x, y, cx, cy (EMU); "keep": "right" / "bottom"
    holds that edge when cx / cy change. The result must stay on the slide and clear of the
    page-number box and the logo."""
    s = ctx.slide(op["sldId"])
    sp = select_sp(s._element, op)
    xf = sp.find(f"{q(P, 'spPr')}/{q(A, 'xfrm')}")
    if xf is None or xf.find(q(A, "off")) is None or xf.find(q(A, "ext")) is None:
        raise EditError(f"resize_shape sldId {s.slide_id}: the shape has no own frame (inherited placeholder)")
    off, ext = xf.find(q(A, "off")), xf.find(q(A, "ext"))
    x, y, cx, cy = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))
    if not any(k in op for k in ("x", "y", "cx", "cy")):
        raise EditError("resize_shape: needs at least one of x, y, cx, cy")
    keep = op.get("keep", [])
    keep = [keep] if isinstance(keep, str) else list(keep)
    for k in keep:
        if k not in ("right", "bottom"):
            raise EditError(f"resize_shape: bad keep {k!r}")
    ncx, ncy = int(op.get("cx", cx)), int(op.get("cy", cy))
    nx = int(op.get("x", x + cx - ncx if "right" in keep else x))
    ny = int(op.get("y", y + cy - ncy if "bottom" in keep else y))
    sw, sh = ctx.prs.slide_width, ctx.prs.slide_height
    if ncx <= 0 or ncy <= 0 or nx < 0 or ny < 0 or nx + ncx > sw or ny + ncy > sh:
        raise EditError(f"resize_shape sldId {s.slide_id}: frame ({nx}, {ny}, {ncx}, {ncy}) leaves the slide")
    for fx0, fy0, fx1, fy1, what, e in _furniture(s._element, sw, sh):
        if e is sp:
            continue
        if min(nx + ncx, fx1) > max(nx, fx0) and min(ny + ncy, fy1) > max(ny, fy0):
            raise EditError(f"resize_shape sldId {s.slide_id}: the new frame overlaps the {what}")
    off.set("x", str(nx))
    off.set("y", str(ny))
    ext.set("cx", str(ncx))
    ext.set("cy", str(ncy))
    ctx.log.append(f"resize_shape sldId {s.slide_id} {sp_name(sp)!r}: ({x}, {y}, {cx}, {cy}) -> ({nx}, {ny}, {ncx}, {ncy})")


def op_set_geometry(ctx, op):
    """Move or resize one picture or shape, selected by "shape" (name) or "shape_id"; any of
    x, y, w, h in inches. The frame must stay on the slide and clear of the page-number box and
    the logo. Works on the XML (not python-pptx's shape cache), so it also works on a slide that
    clone_slide created earlier in the same run."""
    s = ctx.slide(op["sldId"])
    hit = []
    for e in s._element.iter(q(P, "sp"), q(P, "pic"), q(P, "graphicFrame"), q(P, "grpSp"), q(P, "cxnSp")):
        c = next((x for x in e.iter(q(P, "cNvPr"))), None)
        if c is None:
            continue
        if (op.get("shape_id") is not None and c.get("id") == str(op["shape_id"])) or \
           (op.get("shape") is not None and c.get("name") == op["shape"]):
            hit.append(e)
    if len(hit) != 1:
        raise EditError(f"set_geometry sldId {s.slide_id}: selector matched {len(hit)} shapes")
    el = hit[0]
    xf = el.find(f"{q(P, 'spPr')}/{q(A, 'xfrm')}")
    if xf is None:
        xf = el.find(q(P, "xfrm")) if el.tag == q(P, "graphicFrame") else el.find(f"{q(P, 'grpSpPr')}/{q(A, 'xfrm')}")
    if xf is None or xf.find(q(A, "off")) is None or xf.find(q(A, "ext")) is None:
        raise EditError(f"set_geometry sldId {s.slide_id}: the shape has no own frame")
    off, ext = xf.find(q(A, "off")), xf.find(q(A, "ext"))
    if not any(k in op for k in ("x", "y", "w", "h")):
        raise EditError("set_geometry: needs at least one of x, y, w, h")
    cur = {"x": int(off.get("x")), "y": int(off.get("y")), "w": int(ext.get("cx")), "h": int(ext.get("cy"))}
    nx, ny, ncx, ncy = (int(round(op[k] * 914400)) if k in op else cur[k] for k in ("x", "y", "w", "h"))
    sw, shh = ctx.prs.slide_width, ctx.prs.slide_height
    if ncx <= 0 or ncy <= 0 or nx < 0 or ny < 0 or nx + ncx > sw or ny + ncy > shh:
        raise EditError(f"set_geometry sldId {s.slide_id}: frame ({nx}, {ny}, {ncx}, {ncy}) leaves the slide")
    for fx0, fy0, fx1, fy1, what, e in _furniture(s._element, sw, shh):
        if e is el:
            continue
        if min(nx + ncx, fx1) > max(nx, fx0) and min(ny + ncy, fy1) > max(ny, fy0):
            raise EditError(f"set_geometry sldId {s.slide_id}: the new frame overlaps the {what}")
    off.set("x", str(nx))
    off.set("y", str(ny))
    ext.set("cx", str(ncx))
    ext.set("cy", str(ncy))
    name = next(el.iter(q(P, "cNvPr"))).get("name")
    ctx.log.append(f"set_geometry sldId {s.slide_id} {name!r} -> ({nx}, {ny}, {ncx}, {ncy})")


OPS = {
    "replace": op_replace,
    "set_notes": op_set_notes,
    "set_code": op_set_code,
    "set_text": op_set_text,
    "clone_slide": op_clone_slide,
    "replace_picture": op_replace_picture,
    "set_page_number": op_set_page_number,
    "fit_font": op_fit_font,
    "resize_shape": op_resize_shape,
    "cut_code": op_cut_code,
    "set_geometry": op_set_geometry,
    "recolour_code": op_recolour_code,
}


def load_ops(paths):
    ops = []
    for p in paths:
        data = json.loads(Path(p).read_text(encoding="utf-8"))
        if isinstance(data, dict):
            data = data.get("ops", [])
        ops.extend(data)
    return ops


def backup_dir():
    return Path(os.environ.get("DECK_BACKUP_DIR") or Path.home() / ".cache" / "deck_edit" / "backup")


def main(argv):
    ap = argparse.ArgumentParser(description="Apply scripted edits to a .pptx",
                                 usage="deck_edit.py <deck.pptx> <edits.json> [<edits2.json> ...] [--dry-run] [--out PATH]")
    ap.add_argument("deck", type=Path)
    ap.add_argument("edits", type=Path, nargs="+")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--out", type=Path)
    args = ap.parse_args(argv)
    try:
        ops = load_ops(args.edits)
        prs = Presentation(str(args.deck))
        ctx = Ctx(prs)
        for i, op in enumerate(ops, 1):
            fn = OPS.get(op.get("op"))
            if fn is None:
                raise EditError(f"op #{i}: unknown operation {op.get('op')!r}")
            try:
                fn(ctx, op)
            except EditError as e:
                raise EditError(f"op #{i} ({op.get('op')}): {e}")
            except KeyError as e:
                raise EditError(f"op #{i} ({op.get('op')}): missing key {e}")
    except EditError as e:
        print(f"deck_edit: FAILED, nothing written: {e}", file=sys.stderr)
        return 1
    for line in ctx.log:
        print(line)
    target = args.out or args.deck
    tmpdir = Path(tempfile.mkdtemp(prefix="deck_edit_"))
    tmp = tmpdir / target.name
    try:
        prs.save(str(tmp))
        Presentation(str(tmp))  # reload check
        if args.dry_run:
            before, _ = deck_text.pptx_to_md(args.deck)
            after, _ = deck_text.pptx_to_md(tmp)
            diff = list(difflib.unified_diff(before.splitlines(), after.splitlines(), "before", "after", lineterm="", n=1))
            print("\n".join(diff) if diff else "(no text changes)")
            print(json.dumps({"ok": True, "dry_run": True, "created": ctx.created, "ops": len(ops)}))
            return 0
        if args.out is None:
            bdir = backup_dir()
            bdir.mkdir(parents=True, exist_ok=True)
            b = bdir / f"{args.deck.stem}.{time.strftime('%Y%m%d-%H%M%S')}{args.deck.suffix}"
            shutil.copy2(args.deck, b)
            print(f"backup: {b}")
        target.parent.mkdir(parents=True, exist_ok=True)
        os.replace(tmp, target) if tmp.stat().st_dev == (target.parent.stat().st_dev) else shutil.move(str(tmp), str(target))
        print(json.dumps({"ok": True, "out": str(target), "created": ctx.created, "ops": len(ops)}))
        return 0
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
