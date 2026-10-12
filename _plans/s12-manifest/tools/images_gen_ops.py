import json,collections,sys
from pptx import Presentation
F="/tmp/claude-1000/-home-user-Dev/0dea024d-0bf0-4ae5-bddd-7dec0cf60b22/scratchpad/s12/manifest/fork-images"
W=F+"/../work"
idx=json.load(open(F+"/index.json"))
ALL=["Cpp","DotNet","Go","Python","Quarkus","Spring"]
def lang(d,m): return m.get(d)
RULES=[ # (descr prefix, {deck: png}, cat, why)
 ("08b-data-cqrs",{d:"04-cqrs" for d in ALL},"voice","image text had a positional reference ('the next slide') and 'genuinely'"),
 ("08c-data-log-tailing",{d:"04-log-tailing" for d in ALL},"positional","image text had '(next slide)'"),
 ("55-appI-journey",{d:"22-l7-journey" for d in ALL},"voice","subtitle 'each layer earns its place' (ban tier)"),
 ("60-appI-east-west",{d:"22-east-west" for d in ALL},"voice","caption 'for free' (ban tier)"),
 ("86-appK-strangler-shared-database",{d:"24-shared-database" for d in ALL},"positional","'(decomposed DB above)'"),
 ("110-failure-taxonomy",{d:"25-failure-taxonomy" for d in ALL},"voice","'the rest of this appendix' (ban tier)"),
 ("23-anti-patterns-grid",{d:"13-antipatterns-grid" for d in ALL},"voice","box text broken mid-word (layout defect)"),
 ("59-appI-latency",{"Cpp":"22-hop-cost-cpp","Go":"22-hop-cost-go","Python":"22-hop-cost","Quarkus":"22-hop-cost-drools","Spring":"22-hop-cost-drools","DotNet":"22-hop-cost-nrules"},"voice","'every hop must earn its place' (ban tier); rule-engine row per language"),
 ("54-appI-stack",{"Cpp":"22-l7-layers","DotNet":"22-l7-layers","Go":"22-l7-layers"},"stack","in-app router named modern-cpp-kafka (C++), Spring Boot + Camel (in the .NET deck), FastAPI + durable-rules (in the Go deck)"),
 ("61-appI-cpp-router",{"Cpp":"22-business-rule-routing-cpp"},"stack","modern-cpp-kafka did the dispatch; now rdkafka++"),
 ("61-appI-python-router",{"Go":"22-business-rule-routing-go"},"stack","Go deck showed FastAPI + durable-rules; now net/http + gopher-lua + franz-go"),
 ("58-appI-masstransit-routingslip-dn",{"DotNet":"22-business-rule-routing-dotnet"},"stack","MassTransit RoutingSlip removed; NRules + Confluent.Kafka (site Appendix I .NET codetab)"),
 ("03-comm-protocol-matrix",{"Go":"02-pick-protocol-go","DotNet":"02-pick-protocol-dotnet","Spring":"02-pick-protocol-spring"},"stack","framework labels of another language (Python in Go, Quarkus in .NET and Spring)"),
 ("05-comm-service-mesh",{"Go":"02-mesh-handles-wire-go","DotNet":"02-mesh-handles-wire-dotnet","Spring":"02-mesh-handles-wire-spring"},"stack","app box named another language's framework"),
 ("06-composition-federation",{"Go":"03-composition","DotNet":"03-composition","Spring":"03-composition"},"stack","gateway named Strawberry (Go) or SmallRye GraphQL (.NET, Spring)"),
 ("14-stream-topology",{"Go":"06-derive-streams","Python":"06-derive-streams"},"stack","caption 'Faust/Kafka Streams style' (faust-streaming; goka for Go)"),
 ("15-stream-keda-scaling",{"Go":"06-scale-on-lag","DotNet":"06-scale-on-lag","Spring":"06-scale-on-lag"},"stack","HTTP pods named FastAPI (Go) or Quarkus (.NET, Spring)"),
 ("33-appD-state-options",{"Go":"17-saga-state"},"stack","title 'Where a Python saga keeps its state' in the Go deck"),
 ("42-twelve-factor-map-py",{"Python":"01-factors-mapped-py","Go":"01-factors-mapped-go"},"stack","Poetry (examples use pip + requirements.txt); Go deck showed Python tooling"),
 ("117-failure-circuit-breaker.png",{"Go":"25-circuit-breaker"},"stack","library list without the locked Go library (sony/gobreaker)"),
 ("01b-windows-to-k8s-bridge",{"DotNet":"01b-windows-to-k8s-bridge-dn"},"stack","'Temporal workflows / MassTransit sagas'; Serilog row"),
 ("63-appD-stateless-dn",{"DotNet":"63-appD-stateless-dn"},"stack","'or: MassTransit Sagas'; '@Transactional' (Java)"),
 ("70-appJ-pinned-libraries-dn",{"DotNet":"70-appJ-pinned-libraries-dn"},"stack","MassTransit row; 'Every entry above' (positional); rows now from site Appendix J"),
 ("71-appJ-pinned-platform-dn",{"DotNet":"71-appJ-pinned-platform-dn"},"stack","Redis license BSD (Redis 8 is AGPLv3 option); rows from site Appendix J"),
 ("72-appJ-migration-reference-dn",{"DotNet":"72-appJ-migration-reference-dn"},"stack","MassTransit in three rows; rows from site Appendix J"),
 ("29-container-layers",{"DotNet":"29-container-layers"},"engine","old container engine named in a box; separate restore step (S9 changed to one RUN)"),
 ("29-dotnet-compilation",{"DotNet":"29-dotnet-compilation"},"stack","missing-glyph boxes in the old render; site figure"),
 ("29-kestrel-pipeline",{"DotNet":"29-kestrel-pipeline"},"stack","refreshed from the site figure (same content, consistent render)"),
]
CAPTIONS={("Cpp",442):"Drogon is the HTTP front; embedded Lua is the rules engine; rdkafka++ (librdkafka) does the dispatch.",
          ("DotNet",450):"NRules owns the decision; Confluent.Kafka producers and HttpClient dispatch, in-process behind the Istio mesh."}
ops=collections.defaultdict(list); used=collections.Counter(); table=[]
for x in idx:
    for pre,m,cat,why in RULES:
        if x["descr"].startswith(pre) or (pre.endswith(".png") and x["descr"]==pre):
            if pre=="117-failure-circuit-breaker.png" and x["descr"]!=pre: continue
            png=m.get(x["deck"])
            if not png: continue
            ops[x["deck"]].append({"op":"replace_picture","sldId":x["sid"],"shape_id":x["shape_id"],"png":f"lgtm-presentation/edits/png/{png}.png","cat":cat,"_slide":x["n"],"_why":why})
            used[pre]+=1; table.append((x["deck"],x["n"],x["sid"],x["descr"],png))
            break
# captions: find the caption shape (y ~ 6.50in) on that slide
for (deck,sid),text in CAPTIONS.items():
    prs=Presentation(f"~/Dev/cloud-native-design-patterns/lgtm-presentation/Designing-Cloud-Native-APIs-{deck}.pptx")
    s=[t for t in prs.slides if t.slide_id==sid][0]
    cap=[sh for sh in s.shapes if sh.has_text_frame and abs(sh.top/914400-6.50)<0.06]
    assert len(cap)==1,(deck,sid,[sh.name for sh in cap])
    ops[deck].append({"op":"set_text","sldId":sid,"shape_id":cap[0].shape_id,"paragraphs":[text],"cat":"stack","_why":"caption repeated the replaced figure's stale library"})
import os
os.makedirs(W+"/ops-15-images",exist_ok=True)
for d in ALL:
    json.dump({"ops":ops[d]},open(f"{W}/ops-15-images/Designing-Cloud-Native-APIs-{d}.json","w"),indent=1,ensure_ascii=False)
    print(d,len(ops[d]))
json.dump(table,open(F+"/table.json","w"))
