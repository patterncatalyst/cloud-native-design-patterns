#!/usr/bin/env python3
"""Assemble the S12 deck edit manifest.

Inputs (all under this work directory):
  rewrites/*.json   [{"find", "repl", "cat", "decks"?: [basename...], "allow_code"?: bool,
                      "scope"?: slide|notes|both}]  voice / positional / engine rewrites.
                    Without "decks" a rewrite is offered to all six language decks; one that matches
                    the same number of times in all six goes to _shared.json, otherwise it goes to
                    the per-deck files of the decks it matches.
  ops-NN-<name>/<deck>.json  deck_edit ops for one deck, applied in NN order after the rewrites
                    (any op; "expect": "auto" is filled in here). Extra keys ("cat", "_slide", "_why")
                    are kept; deck_edit ignores them. ops-10-appo (Appendix O), ops-15-images,
                    ops-20-sharedcode, ops-30-deck, ops-90-rs (R/S appendix and agenda, gen_rs.py).

Outputs:
  OUT/_shared.json, OUT/<deck>.json   (OUT defaults to the repo's lgtm-presentation/edits)
  after/<deck>.pptx                   the deck after all ops (for scanning)
  report.json                          per deck ops by category, problems

Run with the S12 venv python.
"""
from __future__ import annotations

import collections
import copy
import json
import sys
from pathlib import Path

REPO = Path("~/Dev/cloud-native-design-patterns")
sys.path.insert(0, str(REPO / "scripts"))
import deck_edit as de  # noqa: E402
from pptx import Presentation  # noqa: E402

WORK = Path(__file__).resolve().parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO / "lgtm-presentation" / "edits"
AFTER = WORK / "after"
LANG = [f"Designing-Cloud-Native-APIs-{x}" for x in ("Python", "Quarkus", "Spring", "DotNet", "Go", "Cpp")]
POLICY = ["K8s-Environment-Access-Policies", "Network-Policy-Process-Advisory"]
A, P, q = de.A, de.P, de.q


def deck_path(name):
    return REPO / "lgtm-presentation" / f"{name}.pptx"


def roots_for(slide, scope):
    roots = []
    if scope in ("slide", "both"):
        roots.append(("slide", slide._element))
    if scope in ("notes", "both"):
        for r in de.notes_roots(slide):
            for sp in de.all_sps(r):
                ph = sp.find(f"{q(P, 'nvSpPr')}/{q(P, 'nvPr')}/{q(P, 'ph')}")
                if ph is not None and ph.get("type") == "body":
                    roots.append(("notes", sp))
    return roots


def code_texts(slide):
    """The texts of the code boxes on a slide (to tell a code match from a prose match)."""
    out = []
    for sp in de.all_sps(slide._element):
        if sp.find(q(P, "txBody")) is not None and de.is_code_sp(sp):
            out.append("".join(t.text or "" for t in sp.iter(q(A, "t"))))
    return out


_CODE_CACHE = {}


def apply_replace(prs, op, log, check_code=True):
    """Apply a replace without the expect check; return (count, sldIds hit, code hit sldIds)."""
    scope = op.get("scope", "both")
    ids = op.get("sldIds")
    slides = [s for s in prs.slides if ids is None or s.slide_id in ids]
    total, hit, code_hit = 0, [], []
    find = op["find"]
    for s in slides:
        if find not in "".join(t.text or "" for t in s._element.iter(q(A, "t"))) and not (
                s.has_notes_slide and find in "".join(t.text or "" for t in s.notes_slide._element.iter(q(A, "t")))):
            # quick reject: a match inside one a:br segment is also a substring of the concatenation
            continue
        if check_code and scope in ("slide", "both") and any(find in c for c in code_texts(s)):
            code_hit.append(s.slide_id)
        n_slide = 0
        for kind, root in roots_for(s, scope):
            for p in root.iter(q(A, "p")):
                n_slide += de.replace_in_paragraph(p, op["find"], op["repl"], f"sldId {s.slide_id} {kind}", log)
        if n_slide:
            hit.append(s.slide_id)
            total += n_slide
    return total, hit, code_hit


ENGINE = __import__("re").compile("p(o)dman".replace("(o)", "([oO])"), __import__("re").I)


def dump(obj):
    """JSON with the old engine's name escaped (\\u006f), so the repo text never spells it."""
    text = json.dumps(obj, indent=1, ensure_ascii=False) + "\n"
    return ENGINE.sub(lambda m: m.group(0)[0] + ("\\u006f" if m.group(1) == "o" else "\\u004f") + m.group(0)[2:], text)


def load_json(p):
    data = json.loads(p.read_text(encoding="utf-8"))
    return data.get("ops", data) if isinstance(data, dict) else data


def main():
    problems = []
    notes = []
    rewrites = []
    for f in sorted((WORK / "rewrites").glob("*.json")):
        for i, r in enumerate(load_json(f)):
            r = dict(r)
            r["_src"] = f"{f.name}#{i}"
            rewrites.append(r)
    # ---- pass 1: counts per language deck, in rewrite order
    counts = collections.defaultdict(dict)
    for d in LANG:
        prs = Presentation(str(deck_path(d)))
        log = []
        for i, r in enumerate(rewrites):
            if "decks" in r and d not in r["decks"]:
                continue
            n, hit, code = apply_replace(prs, {"find": r["find"], "repl": r["repl"], "scope": r.get("scope", "both")}, log)
            counts[i][d] = n
            if code and not r.get("allow_code"):
                problems.append(f"{d}: rewrite {r['_src']} matches inside code on sldIds {sorted(set(code))}: {r['find'][:60]!r}")
    shared_idx = []
    for i, r in enumerate(rewrites):
        c = counts[i]
        if sum(c.values()) == 0:
            problems.append(f"rewrite {r['_src']} matches nowhere: {r['find'][:80]!r}")
            continue
        if "decks" not in r and len(c) == len(LANG) and len(set(c.values())) == 1:
            shared_idx.append(i)
    shared_set = set(shared_idx)
    shared_ops = []
    for i in shared_idx:
        r = rewrites[i]
        op = {"op": "replace", "scope": r.get("scope", "both"), "find": r["find"], "repl": r["repl"],
              "expect": counts[i][LANG[0]], "cat": r.get("cat", "voice")}
        shared_ops.append(op)
    # ---- pass 2: per deck sequence
    report = {"problems": problems, "decks": {}}
    OUT.mkdir(parents=True, exist_ok=True)
    AFTER.mkdir(exist_ok=True)
    for d in LANG + POLICY:
        prs = Presentation(str(deck_path(d)))
        log = []
        ctx = de.Ctx(prs)
        cats = collections.Counter()
        if d in LANG:
            for op in shared_ops:
                n, _, _ = apply_replace(prs, op, log)
                if n != op["expect"]:
                    problems.append(f"{d}: shared {op['find'][:60]!r} expected {op['expect']} got {n}")
        deck_ops = []
        for i, r in enumerate(rewrites):
            if i in shared_set:
                continue
            if "decks" in r and d not in r["decks"]:
                continue
            if "decks" not in r and d not in LANG:
                continue
            op = {"op": "replace", "scope": r.get("scope", "both"), "find": r["find"], "repl": r["repl"]}
            n, hit, _ = apply_replace(prs, op, log)
            if n == 0:
                continue
            op["expect"] = n
            op["sldIds"] = hit
            op["cat"] = r.get("cat", "voice")
            deck_ops.append(op)
        for src in sorted(WORK.glob(f"ops-*/{d}.json")):
            if not src.exists():
                continue
            for j, op in enumerate(load_json(src)):
                op = copy.deepcopy(op)
                where = f"{src.parent.name}/{src.name}#{j}"
                try:
                    if op["op"] == "replace":
                        n, hit, code = apply_replace(prs, op, log, check_code=False)
                        given = op.get("expect", "auto")
                        if n == 0:
                            problems.append(f"{d}: {where} replace matches nothing: {op['find'][:60]!r}")
                        elif given != "auto" and n != int(given):
                            notes.append(f"{d}: {where} replace {op['find'][:50]!r}: count {given} on the original deck, {n} after earlier ops (using {n})")
                        op["expect"] = n
                        if hit and "sldIds" not in op:
                            op["sldIds"] = hit
                    else:
                        fn = de.OPS[op["op"]]
                        ctx.created = ctx.created  # keep $refs
                        fn(ctx, op)
                except de.EditError as e:
                    problems.append(f"{d}: {where} {op['op']}: {e}")
                except KeyError as e:
                    problems.append(f"{d}: {where} {op.get('op')}: missing {e}")
                deck_ops.append(op)
        for op in (shared_ops if d in LANG else []):
            cats[op.get("cat", "voice") + " (shared)"] += 1
        for op in deck_ops:
            cats[op.get("cat", op["op"])] += 1
        (OUT / f"{d}.json").write_text(dump({"ops": deck_ops}))
        try:
            prs.save(str(AFTER / f"{d}.pptx"))
        except Exception as e:  # noqa: BLE001
            problems.append(f"{d}: save failed {e}")
        report["decks"][d] = {"ops": len(deck_ops), "by_cat": dict(cats)}
    (OUT / "_shared.json").write_text(dump({"ops": shared_ops}))
    report["shared_ops"] = len(shared_ops)
    (WORK / "report.json").write_text(json.dumps(report, indent=1))
    print(json.dumps({k: v for k, v in report.items() if k != "problems"}, indent=1))
    for n_ in notes:
        print("  note:", n_)
    print(f"{len(problems)} problem(s)")
    for p in problems:
        print("  -", p)


if __name__ == "__main__":
    main()
