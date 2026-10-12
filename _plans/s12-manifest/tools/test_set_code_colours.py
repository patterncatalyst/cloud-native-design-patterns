#!/usr/bin/env python3
"""Regression test for deck_edit.py set_code / recolour_code colours (S12 fix round A).

The bug: code_lines() took the run that ended the previous line ("# comment\\n") as the template
of the next line, so set_code wrote code lines in the comment green 8FB98F. This test builds
small decks with code boxes in the three layouts the decks use (newlines embedded in one
paragraph, one paragraph per line, a:br breaks), each starting with a comment line, runs
set_code and recolour_code on them, and checks every run colour by line type.

    <venv>/bin/python _plans/s12-manifest/tools/test_set_code_colours.py    # exit 0 = pass
"""
from __future__ import annotations

import copy
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import deck_edit as de  # noqa: E402
import colour_audit  # noqa: E402
from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

A, P = de.A, de.P
GREEN, WHITE = de.COMMENT_RGB, de.CODE_RGB


def run(text, rgb):
    r = etree.Element(de.q(A, "r"))
    rpr = etree.SubElement(r, de.q(A, "rPr"), lang="en-US", sz="1200")
    sf = etree.SubElement(rpr, de.q(A, "solidFill"))
    etree.SubElement(sf, de.q(A, "srgbClr"), val=rgb)
    etree.SubElement(rpr, de.q(A, "latin"), typeface="Red Hat Mono")
    etree.SubElement(r, de.q(A, "t")).text = text
    return r


def br(rgb):
    b = etree.Element(de.q(A, "br"))
    b.append(copy.deepcopy(run("", rgb).find(de.q(A, "rPr"))))
    return b


def add_box(slide, name, paragraphs):
    """paragraphs: list of lists of elements (runs / breaks)."""
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(1), Inches(12), Inches(5.5))
    tb.name = name
    tx = tb._element.find(de.q(P, "txBody"))
    for p in tx.findall(de.q(A, "p")):
        tx.remove(p)
    for elems in paragraphs:
        p = etree.SubElement(tx, de.q(A, "p"))
        for e in elems:
            p.append(e)
    return tb


def build(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout = prs.slide_layouts[6]
    # 1: embedded newlines, one paragraph; the comment run ends with "\n" (the original bug)
    s1 = prs.slides.add_slide(layout)
    add_box(s1, "embedded", [[run("# a comment first\n", GREEN), run("x = 1\ny = 2", WHITE)]])
    # 2: one paragraph per line, first paragraph a comment
    s2 = prs.slides.add_slide(layout)
    add_box(s2, "paragraphs", [[run("// a comment first", GREEN)], [run("int x = 1;", WHITE)],
                               [run("int y = 2;", WHITE)]])
    # 3: a:br breaks, first line a comment
    s3 = prs.slides.add_slide(layout)
    add_box(s3, "breaks", [[run("// a comment first", GREEN), br(GREEN), run("int x = 1;", WHITE),
                            br(WHITE), run("int y = 2;", WHITE)]])
    # 4: a box with no comment line at all (comments must still come out green)
    s4 = prs.slides.add_slide(layout)
    add_box(s4, "nocomment", [[run("a = 1\nb = 2\nc = 3", WHITE)]])
    add_box(s4, "nocomment-paras", [[run("a = 1", WHITE)], [run("b = 2", WHITE)], [run("c = 3", WHITE)]])
    # 5: a box already miscoloured (code lines green): recolour_code repairs it
    s5 = prs.slides.add_slide(layout)
    add_box(s5, "miscoloured", [[run("# head\n", GREEN), run("x = 1\n", GREEN), run("y = 2  ", GREEN),
                                 run("# trailing", GREEN)]])
    prs.save(path)
    return [s.slide_id for s in prs.slides]


NEW = {
    "embedded": ["# rewritten comment", "first = call()", "second = other()  # inline", "",
                 "# another comment", "third = 3"],
    "paragraphs": ["// rewritten comment", "int a = 1;", "", "int b = 2;  // inline"],
    "breaks": ["// rewritten comment", "int a = 1;", "// middle comment", "int b = 2;"],
    "nocomment": ["# a comment in a box that had none", "a = 10", "b = 20"],
}


def main():
    fails = []
    with tempfile.TemporaryDirectory() as td:
        src, out = Path(td) / "t.pptx", Path(td) / "out.pptx"
        ids = build(src)
        prs = Presentation(str(src))
        ctx = de.Ctx(prs)
        for sid, name in zip(ids, ["embedded", "paragraphs", "breaks", "nocomment"]):
            de.op_set_code(ctx, {"op": "set_code", "sldId": sid, "shape": name, "lines": NEW[name]})
        de.op_recolour_code(ctx, {"op": "recolour_code", "sldId": ids[4], "shape": "miscoloured"})
        try:
            de.op_recolour_code(ctx, {"op": "recolour_code", "sldId": ids[4], "shape": "miscoloured"})
            fails.append("recolour_code on an already-correct box did not fail without allow_noop")
        except de.EditError:
            pass
        prs.save(str(out))
        hits = colour_audit.audit(out)
        for h in hits:
            fails.append(f"slide {h['slide']} line {h['line']} {h['kind']}: {h['text']!r} {h['colours']}")
        # explicit check of the original bug: the first code line after a leading comment is E6E6E6
        prs2 = Presentation(str(out))
        for n, (s, name) in enumerate(zip(prs2.slides, ["embedded", "paragraphs", "breaks", "nocomment",
                                                        "miscoloured"]), 1):
            sp = next(x for x in de.all_sps(s._element) if de.sp_name(x) == name)
            ls = colour_audit.box_lines(sp.find(de.q(P, "txBody")))
            texts = ["".join(t for t, _ in l) for l in ls]
            want = NEW.get(name, ["# head", "x = 1", "y = 2  # trailing"])
            if texts != want:
                fails.append(f"{name}: text {texts!r} != {want!r}")
            code1 = next(l for l in ls if "".join(t for t, _ in l).strip()
                         and not de.is_comment("".join(t for t, _ in l)))
            if [c for t, c in code1 if t.strip()][0] != WHITE:
                fails.append(f"{name}: first code line {code1!r} is not {WHITE}")
            com = next(l for l in ls if de.is_comment("".join(t for t, _ in l)))
            if any(c != GREEN for t, c in com if t.strip()):
                fails.append(f"{name}: comment line {com!r} is not {GREEN}")
        # the template fix on its own: code_lines picks the code run, not the "...\n" comment run
        prs3 = Presentation(str(src))
        sp = next(x for x in de.all_sps(prs3.slides[0]._element) if de.sp_name(x) == "embedded")
        ls = de.code_lines(sp.find(de.q(P, "txBody")))
        if de.run_rgb(ls[1]["run"]) != WHITE:
            fails.append(f"code_lines: line 2 template run is {de.run_rgb(ls[1]['run'])}, want {WHITE}")
        # cut_code: an ellipsis in a box without a comment line is comment-coloured
        prs5 = Presentation(str(src))
        ctx5 = de.Ctx(prs5)
        de.op_cut_code(ctx5, {"op": "cut_code", "sldId": ids[3], "shape": "nocomment-paras",
                              "cuts": [{"from": 2, "to": 2, "first": "b = 2", "ellipsis": "# ... b elided"}]})
        sp = next(x for x in de.all_sps(prs5.slides[3]._element) if de.sp_name(x) == "nocomment-paras")
        ell = [l for l in colour_audit.box_lines(sp.find(de.q(P, "txBody")))
               if "".join(t for t, _ in l).startswith("# ...")]
        if not ell or any(c != GREEN for t, c in ell[0] if t.strip()):
            fails.append(f"cut_code: ellipsis line {ell!r} is not {GREEN}")
        # colour: false keeps the template colours (no recolouring)
        prs4 = Presentation(str(src))
        ctx4 = de.Ctx(prs4)
        de.op_set_code(ctx4, {"op": "set_code", "sldId": ids[3], "shape": "nocomment", "colour": False,
                              "lines": NEW["nocomment"]})
        if "recoloured" in ctx4.log[-1]:
            fails.append("colour: false still recoloured")
    for f in fails:
        print("FAIL", f)
    print("test_set_code_colours:", "FAILED" if fails else "ok", f"({len(fails)} failure(s))")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
