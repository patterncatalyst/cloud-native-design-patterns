"""Extract every text segment (paragraph split at a:br) of each deck, with sldId, kind, shape, code flag."""
import json, sys
from pathlib import Path
sys.path.insert(0, "~/Dev/cloud-native-design-patterns/scripts")
import deck_edit as de
from pptx import Presentation
A=de.A; P=de.P; q=de.q
ROOT=Path("~/Dev/cloud-native-design-patterns/lgtm-presentation")
out={}
for deck in sorted(ROOT.glob("*.pptx")):
    prs=Presentation(str(deck))
    segs=[]
    for n,s in enumerate(prs.slides,1):
        roots=[("slide",s._element)]
        if s.has_notes_slide:
            for sp in de.all_sps(s.notes_slide._element):
                ph=sp.find(f"{q(P,'nvSpPr')}/{q(P,'nvPr')}/{q(P,'ph')}")
                if ph is not None and ph.get("type")=="body":
                    roots.append(("notes",sp))
        for kind,root in roots:
            sps = de.all_sps(root) if kind=="slide" else [root]
            for sp in sps:
                code = de.is_code_sp(sp) if kind=="slide" else False
                name = de.sp_name(sp)
                for p in sp.iter(q(A,"p")):
                    cur=""
                    parts=[]
                    for c in p:
                        if c.tag==q(A,"r"): cur+=c.findtext(q(A,"t")) or ""
                        elif c.tag==q(A,"br"): parts.append(cur); cur=""
                    parts.append(cur)
                    for t in parts:
                        if t.strip():
                            segs.append(dict(n=n,sid=s.slide_id,kind=kind,shape=name,code=code,text=t))
            # tables
        for gf in s._element.iter(q(P,"graphicFrame")):
            for tc in gf.iter(q(A,"tc")):
                for p in tc.iter(q(A,"p")):
                    t=de.ptext(p)
                    if t.strip(): segs.append(dict(n=n,sid=s.slide_id,kind="slide",shape="table",code=False,text=t))
    out[deck.stem]=segs
    print(deck.stem,len(prs.slides),len(segs))
Path("segs.json").write_text(json.dumps(out))
