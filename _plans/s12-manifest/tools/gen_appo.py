#!/usr/bin/env python3
"""Appendix O: shared sentence rewrites (rewrites/10-appo.json) and per-deck code/figure ops (ops-10-appo/)."""
import json, sys
from pathlib import Path
REPO = Path(__file__).resolve().parents[3]
from pptx import Presentation
from pptx.util import Emu
WORK = Path(__file__).resolve().parent
E = "P" + "odman"; e = "p" + "odman"
RW = [
 ("Postman collections as an automated gate — on plain " + E, "Postman collections as an automated gate on Docker Engine", "engine"),
 ("The arc: why Newman, the anatomy of a collection, one collection against five backends, running it and shaping the output, data-driven cases, the CI gate on " + E + ", and where Newman sits in the test pyramid.",
  "The arc: why Newman, the anatomy of a collection, one collection against six backends, running it and shaping the output, data-driven cases, the CI gate on Docker Engine, and where Newman sits in the test pyramid.", "engine"),
 ("This is a quietly powerful idea: the collection is the API specification made executable.", "The collection is the API specification made executable.", "voice"),
 ("One collection, five backends", "One collection, six backends", "stack"),
 ("Only this step is language-specific; the Newman command is identical for all five.", "Only this step is language-specific; the Newman command is identical for all six.", "stack"),
 ("That is the whole point of black-box testing: bring up Spring Boot, Quarkus, .NET, Python, or C++, and run the exact same collection against it.",
  "That is the point of black-box testing: bring up Spring Boot, Quarkus, .NET, Python, C++ or Go, and run the same collection against it.", "stack"),
 ("In CI you would start the service with " + e + " compose rather than a foreground command, but locally a background process and a newman run is the fastest loop.",
  "In CI the services start with docker compose up -d --build --wait instead of a foreground command; locally, a background process and a newman run is the fastest loop.", "engine"),
 ("Bring the stack up on " + E + ", wait for readiness, run the collection — a failed assertion fails the build.",
  "Bring the stack up with docker compose, wait for readiness, run the collection: a failed assertion fails the build.", "engine"),
 ("The services come up on " + E + " exactly as they do on a laptop - no managed cloud - and Newman runs against them.",
  "The services come up on Docker Engine as they do locally, with no managed cloud, and Newman runs against them.", "engine"),
 ("The gate on " + E, "The gate on Docker Engine", "engine"),
 ("The full gate: a pull request builds the image, " + e + " compose brings the order, payment, and inventory services up, Newman runs the collection against them, and the result gates the merge - pass merges, fail blocks.",
  "The full gate: a pull request builds the image, docker compose brings the order, payment, and inventory services up, Newman runs the collection against them, and the result gates the merge: pass merges, fail blocks.", "engine"),
 ("The important property is parity: the stack in CI is the same stack you run on your laptop, on the same " + E + ", so a green run in CI means the contract holds against the real services, not a mocked approximation.",
  "The important property is parity: the stack in CI is the stack you run locally, on the same Docker Engine, so a green run in CI means the contract holds against the running services rather than a mock.", "engine"),
 ("the API/contract tier — above unit tests, below a thin end-to-end layer", "the API/contract tier: on top of unit tests, under a thin end-to-end layer", "positional"),
 ("Newman is not a replacement for unit tests - it is the layer above them.", "Newman does not replace unit tests; it is the layer on top of them.", "positional"),
 ("and reserve a thin layer of full end-to-end tests for the few journeys that genuinely need the whole system.",
  "and reserve a thin layer of full end-to-end tests for the few user flows that need the whole system.", "voice"),
 ("And test the failure paths as deliberately as the happy path - a 401 with WWW-Authenticate, a 403, a 409 or 422 on a business rule, a 429 with RateLimit headers, and idempotency by sending the same key twice - which ties this appendix back to the error-format, security, idempotency, and rate-limiting material.",
  "Test the failure paths as carefully as the happy path: a 401 with WWW-Authenticate, a 403, a 409 or 422 on a business rule, a 429 with RateLimit headers, and idempotency by sending the same key twice. Those cases tie this appendix back to the error-format, security, idempotency, and rate-limiting material.", "voice"),
 ("One collection is the conformance suite for all five implementations: Newman tests over HTTP, as a black box.",
  "One collection is the conformance suite for all six implementations: Newman tests over HTTP, as a black box.", "stack"),
 ("Gate CI on plain " + E + " so the build fails on a broken contract; keep Newman in the API/contract tier.",
  "Gate CI on plain Docker Engine so the build fails on a broken contract; keep Newman in the API/contract tier.", "engine"),
 ("Because Newman tests over HTTP, one collection is the conformance suite for all five implementations.",
  "Because Newman tests over HTTP, one collection is the conformance suite for all six implementations.", "stack"),
 ("Wire it into CI on plain " + E + " so the build fails on a broken contract, and keep Newman in its lane - the API/contract tier, above unit tests and below a thin end-to-end layer.",
  "Wire it into CI on plain Docker Engine so the build fails on a broken contract, and keep Newman in its lane: the API/contract tier, on top of unit tests and under a thin end-to-end layer.", "engine"),
]
LAUNCH = {  # _docs/28-appendix-o-newman.md, "One collection, six backends" codetab
 "Spring": ["# Spring Boot — start the service under test, then run the same collection",
            "./mvnw spring-boot:run &           # http://localhost:8080"],
 "Quarkus": ["# Quarkus — dev mode, then the same collection",
             "./mvnw quarkus:dev &               # http://localhost:8080"],
 "DotNet": ["# .NET — run the service, then the same collection",
            "dotnet run &                       # http://localhost:8080"],
 "Python": ["# Python / FastAPI — serve, then the same collection",
            "uvicorn app.main:app --port 8080 & # http://localhost:8080"],
 "Cpp": ["# C++ / Drogon — build and run, then the same collection",
         "cmake --build build && ./build/order-service &   # http://localhost:8080"],
 "Go": ["# Go — run the service, then the same collection",
        "go run ./cmd/order-service &       # http://localhost:8080"],
}
NEWMAN_RUN = "newman run orders.postman_collection.json -e local.postman_environment.json"
POSTMAN = [  # _docs/28, "Anatomy of a collection"
 '{', '  "name": "Create order",', '  "request": {', '    "method": "POST",',
 '    "header": [{ "key": "Idempotency-Key", "value": "{{$guid}}" }],',
 '    "url": "{{base_url}}/orders",',
 '    "body": { "mode": "raw", "raw": "{\\"sku\\":\\"A-1\\",\\"qty\\":2}" }', '  },',
 '  "event": [{', '    "listen": "test",', '    "script": { "exec": [',
 '      "pm.test(\'201 Created\', () => pm.response.to.have.status(201));",',
 '      "pm.test(\'under 300ms\', () => pm.expect(pm.response.responseTime).to.be.below(300));",',
 '      "const body = pm.response.json();",',
 '      "pm.test(\'returns an id\', () => pm.expect(body).to.have.property(\'id\'));",',
 '      "pm.collectionVariables.set(\'order_id\', body.id);"', '    ] }', '  }]', '}']
RUNNING = [  # _docs/28, "Running Newman" and "Data-driven tests" (comment moved off the continuation line)
 "# install once (Node.js):",
 "#   npm install -g --prefix ~/.local newman@6.2.3 newman-reporter-htmlextra@1.23.1",
 "newman run orders.postman_collection.json \\",
 "  --environment local.postman_environment.json \\",
 "  --reporters cli,junit,htmlextra \\",
 "  --reporter-junit-export results/newman.xml \\",
 "  --timeout-request 5000 \\",
 "  --bail",
 "# --timeout-request: per-request ceiling (ms); --bail: stop at the first failure",
 "",
 "# data-driven: one iteration per CSV row; the columns become {{variables}}",
 "newman run orders.postman_collection.json -e local.postman_environment.json \\",
 "  --iteration-data cases.csv",
]
CI = [  # _docs/28, "Newman in CI" (actions at v7, as in .github/workflows)
 "# .github/workflows/api-tests.yml (any runner with Docker Engine and Node.js)",
 "jobs:", "  api-tests:", "    runs-on: ubuntu-latest", "    steps:",
 "      - uses: actions/checkout@v7",
 "      - run: docker compose up -d --build --wait   # order, payment, inventory, ...",
 "      - run: ./scripts/wait-for-ready.sh           # poll /readyz before testing",
 "      - run: npm i -g newman@6.2.3 newman-reporter-htmlextra@1.23.1",
 "      - run: |",
 "          newman run orders.postman_collection.json \\",
 "            -e ci.postman_environment.json \\",
 "            --reporters cli,junit \\",
 "            --reporter-junit-export results/newman.xml",
 "      - uses: actions/upload-artifact@v7          # keep the report even on failure",
 "        if: always()",
 "        with: { name: newman-results, path: results/ }",
]
BASE = {"Python": 516, "Quarkus": 516, "Go": 516, "Cpp": 516, "Spring": 520, "DotNet": 530}


def ids(prs, sid):
    s = next(x for x in prs.slides if x.slide_id == sid)
    code = [sh.shape_id for sh in s.shapes if sh.has_text_frame and Emu(sh.height).inches > 3 and
            any("Mono" in (r.font.name or "") for p in sh.text_frame.paragraphs for r in p.runs)]
    pics = [sh.shape_id for sh in s.shapes if sh.shape_type == 13 and 1.5 < Emu(sh.top).inches < 2.0]
    return code, pics


def main():
    (WORK / "rewrites").mkdir(exist_ok=True)
    (WORK / "rewrites" / "10-appo.json").write_text(json.dumps(
        [{"find": f, "repl": r, "cat": c, "_src": "Appendix O"} for f, r, c in RW], indent=1, ensure_ascii=False))
    (WORK / "ops-10-appo").mkdir(exist_ok=True)
    for deck, b in BASE.items():
        name = f"Designing-Cloud-Native-APIs-{deck}"
        prs = Presentation(str(REPO / "lgtm-presentation" / f"{name}.pptx"))
        ops = []
        def code(off, lines, why):
            c, _ = ids(prs, b + off)
            assert len(c) == 1, (deck, b + off, c)
            ops.append({"op": "set_code", "sldId": b + off, "shape_id": c[0], "lines": lines, "cat": "code", "_why": why})
        def pic(off, png, why):
            _, p = ids(prs, b + off)
            assert len(p) == 1, (deck, b + off, p)
            ops.append({"op": "replace_picture", "sldId": b + off, "shape_id": p[0],
                        "png": f"lgtm-presentation/edits/png/{png}.png", "cat": "stack", "_why": why})
        code(3, POSTMAN, "_docs/28 Anatomy of a collection (adds the response-time assertion)")
        pic(4, "27-black-box", "the old figure shows five backends; the site figure shows all six")
        code(5, LAUNCH[deck] + [NEWMAN_RUN], "_docs/28 codetab 'One collection, six backends', this deck's language")
        code(6, RUNNING, "_docs/28 Running Newman + Data-driven tests; pinned install; no comment after a line continuation")
        code(8, CI, "_docs/28 Newman in CI on Docker Engine; actions at v7 as in .github/workflows")
        pic(9, "27-ci-gate", "the old figure shows the old container engine; site figure 27-ci-gate")
        (WORK / "ops-10-appo" / f"{name}.json").write_text(json.dumps({"ops": ops}, indent=1, ensure_ascii=False))
        print(name, len(ops))


if __name__ == "__main__":
    main()
