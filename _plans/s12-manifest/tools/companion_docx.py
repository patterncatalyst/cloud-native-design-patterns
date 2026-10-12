#!/usr/bin/env python3
"""Edit a policy companion .docx in place so its text equals a new revision of its .md (decision 7).

    companion_docx.py <old.md> <new.md> <in.docx> <out.docx> [--ops ops.json] [--png rIdN=path ...]

The docx was built from the md with one paragraph per md line (headings, bullets, numbered items,
continuation lines, one paragraph per fenced block, two per image: the picture and its alt text) and a
table per md table. This script parses both md revisions into that paragraph model, checks that the old md
reproduces the docx text exactly, aligns old and new with difflib, and applies the difference as
paragraph-level operations: set the text of a paragraph (runs rebuilt from the md's **bold**, `code` and
*italic* spans, formatting copied from existing runs of the same kind), insert a paragraph (a clone of an
existing paragraph of the same style) or delete one. Tables keep their shape; changed cells get new text.
--png replaces an embedded image by relationship id and rescales its height to the new aspect ratio.
At the end the output docx is re-read and compared with the new md; any difference is an error.
Run with the S12 venv python (python-docx 1.2.0).
"""
import copy
import difflib
import json
import re
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

LINK = re.compile(r"\[([^\]]+)\]\([^)]+\)")


def spans(text):
    """md inline -> [(text, kind)]; kind in plain, bold, code, italic (as the original converter did)."""
    text = LINK.sub(r"\1", text)
    out = []
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if i % 2:
            out.append((part, "bold"))
            continue
        for j, p2 in enumerate(re.split(r"`([^`]+)`", part)):
            if j % 2:
                out.append((p2, "code"))
                continue
            for k, p3 in enumerate(re.split(r"(?<![\w*])\*([^*\s][^*]*?)\*(?![\w*])", p2)):
                if p3:
                    out.append((p3, "italic" if k % 2 else "plain"))
    return [s for s in out if s[0]]


def parse(md):
    blocks, lines, i = [], md.split("\n"), 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            j = i + 1
            while not lines[j].startswith("```"):
                j += 1
            blocks.append({"style": "Normal", "spans": [("\n".join(lines[i + 1:j]), "plain")], "code": True})
            i = j + 1
            continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in cells):
                    rows.append([[(c, "plain")] for c in cells])  # cells keep their md text
                i += 1
            blocks.append({"table": rows})
            continue
        i += 1
        if not ln.strip() or ln.strip() == "---":
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", ln)
        if m:
            blocks.append({"style": "Normal", "spans": [], "image": m.group(2)})
            blocks.append({"style": "Normal", "spans": [(m.group(1), "plain")]})
            continue
        for rx, style in ((r"^# (.*)$", "Heading 1"), (r"^## (.*)$", "Heading 2"), (r"^### (.*)$", "Heading 3"),
                          (r"^- \[ \] (.*)$", "List Bullet"), (r"^\s*- (.*)$", "List Bullet"),
                          (r"^\d+\. (.*)$", "List Number")):
            m = re.match(rx, ln)
            if m:
                blocks.append({"style": style, "spans": spans(m.group(1))})
                break
        else:
            blocks.append({"style": "Normal", "spans": spans(ln)})
    return blocks


def btext(b):
    if "table" in b:
        return "TABLE:" + json.dumps([["".join(t for t, _ in c) for c in r] for r in b["table"]])
    return b["style"] + ":" + "".join(t for t, _ in b["spans"])


def body(doc):
    out = []
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            out.append(Paragraph(el, doc))
        elif el.tag == qn("w:tbl"):
            out.append(Table(el, doc))
    return out


def dtext(x):
    if isinstance(x, Table):
        return "TABLE:" + json.dumps([[c.text for c in r.cells] for r in x.rows])
    return x.style.name + ":" + x.text


class Runs:
    def __init__(self, doc):
        self.t = {}
        for p in doc.paragraphs:
            for r in p.runs:
                k = "code" if (r.font.name or "").startswith("Consolas") else "bold" if r.bold else \
                    "italic" if r.italic else "plain"
                self.t.setdefault((p.style.name, k), r._r)
                self.t.setdefault(("*", k), r._r)

    def fill(self, p, sp, style):
        own = {}
        for r in p.runs:
            k = "code" if (r.font.name or "").startswith("Consolas") else "bold" if r.bold else \
                "italic" if r.italic else "plain"
            own.setdefault(k, r._r)
        for r in list(p._p.iterchildren(qn("w:r"))):
            p._p.remove(r)
        for text, kind in sp:
            src = next(x for x in (own.get(kind), self.t.get((style, kind)), self.t.get(("*", kind)),
                                   self.t[("*", "plain")]) if x is not None)
            r = copy.deepcopy(src)
            for c in list(r):
                if c.tag != qn("w:rPr"):
                    r.remove(c)
            parts = text.split("\n")
            for n, piece in enumerate(parts):
                if n:
                    r.append(r.makeelement(qn("w:br"), {}))
                t = r.makeelement(qn("w:t"), {})
                t.text = piece
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                r.append(t)
            p._p.append(r)


def main(argv):
    old_md, new_md, src, dst = argv[1:5]
    ops_path = argv[argv.index("--ops") + 1] if "--ops" in argv else None
    pngs = [a.split("=", 1) for a in argv[argv.index("--png") + 1:]] if "--png" in argv else []
    old, new = parse(open(old_md).read()), parse(open(new_md).read())
    doc = Document(src)
    items = body(doc)
    got = [dtext(x) for x in items]
    want_old = [btext(b) for b in old]
    if got != want_old:
        for a, b in zip(got, want_old):
            if a != b:
                raise SystemExit(f"old md does not reproduce the docx:\n  docx: {a[:120]}\n  md:   {b[:120]}")
        raise SystemExit(f"old md does not reproduce the docx: {len(got)} vs {len(want_old)} blocks")
    runs = Runs(doc)
    styles = {}
    for x in items:
        if isinstance(x, Paragraph):
            styles.setdefault(x.style.name, x)
    ops = []
    sm = difflib.SequenceMatcher(a=want_old, b=[btext(b) for b in new], autojunk=False)
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        n = max(i2 - i1, j2 - j1)
        anchor = items[i1 - 1] if i1 else None
        for k in range(n):
            i, j = i1 + k, j1 + k
            nb = new[j] if j < j2 else None
            if i < i2 and nb is not None:
                x = items[i]
                if isinstance(x, Table) or "table" in nb:
                    if not (isinstance(x, Table) and "table" in nb):
                        raise SystemExit(f"table/paragraph mismatch at old block {i}")
                    for r, row in zip(x.rows, nb["table"]):
                        for c, sp in zip(r.cells, row):
                            if c.text != "".join(t for t, _ in sp):
                                runs.fill(c.paragraphs[0], sp, c.paragraphs[0].style.name)
                                ops.append({"op": "set_cell", "block": i, "text": "".join(t for t, _ in sp)})
                    anchor = x
                    continue
                if x.style.name != nb["style"]:
                    x.style = doc.styles[nb["style"]]
                runs.fill(x, nb["spans"], nb["style"])
                ops.append({"op": "set", "block": i, "style": nb["style"], "text": "".join(t for t, _ in nb["spans"])})
                anchor = x
            elif nb is not None:
                if "table" in nb or nb.get("image") is not None:
                    raise SystemExit(f"cannot insert a table or image (new block {j})")
                tmpl = styles[nb["style"]]
                el = copy.deepcopy(tmpl._p)
                if anchor is None:
                    raise SystemExit("insert before the first block")
                (anchor._element if isinstance(anchor, Table) else anchor._p).addnext(el)
                p = Paragraph(el, doc)
                runs.fill(p, nb["spans"], nb["style"])
                ops.append({"op": "insert_after", "after_text": dtext(anchor)[:80], "style": nb["style"],
                            "text": "".join(t for t, _ in nb["spans"])})
                anchor = p
            else:
                x = items[i]
                (x._element if isinstance(x, Table) else x._p).getparent().remove(
                    x._element if isinstance(x, Table) else x._p)
                ops.append({"op": "delete", "block": i, "text": dtext(x)})
    for rid, path in pngs:
        part = doc.part.related_parts[rid]
        from PIL import Image
        w, h = Image.open(path).size
        part._blob = open(path, "rb").read()
        for blip in doc.element.body.iter(qn("a:blip")):
            if blip.get(qn("r:embed")) == rid:
                inline = blip
                while inline.tag not in (qn("wp:inline"), qn("wp:anchor")):
                    inline = inline.getparent()
                ext = inline.find(qn("wp:extent"))
                cx = int(ext.get("cx"))
                ext.set("cy", str(round(cx * h / w)))
                for e in inline.iter(qn("a:ext")):
                    e.set("cx", str(cx))
                    e.set("cy", str(round(cx * h / w)))
        ops.append({"op": "replace_image", "rId": rid, "png": path, "size": [w, h]})
    doc.save(dst)
    check = [dtext(x) for x in body(Document(dst))]
    want = [btext(b) for b in new]
    if check != want:
        for a, b in zip(check, want):
            if a != b:
                raise SystemExit(f"result differs from the new md:\n  docx: {a[:120]}\n  md:   {b[:120]}")
        raise SystemExit(f"result differs from the new md: {len(check)} vs {len(want)} blocks")
    if ops_path:
        open(ops_path, "w").write(json.dumps(ops, indent=1, ensure_ascii=False) + "\n")
    print(json.dumps({"ok": True, "set": sum(o["op"] == "set" for o in ops),
                      "set_cell": sum(o["op"] == "set_cell" for o in ops),
                      "insert": sum(o["op"] == "insert_after" for o in ops),
                      "delete": sum(o["op"] == "delete" for o in ops),
                      "images": sum(o["op"] == "replace_image" for o in ops), "blocks": len(want)}))


if __name__ == "__main__":
    main(sys.argv)
